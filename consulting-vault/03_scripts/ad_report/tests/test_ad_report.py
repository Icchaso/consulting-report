"""広告レポート（ad_analyze / build_deck / check_deck / make_sample）のテスト。

一時ディレクトリに vault を組み立てて、各スクリプトを --vault 付きで実行して確かめる。
実際の vault の 01_clients/ には何も書かない。

実行（vault ルートで）:
  python3 -m unittest discover -s 03_scripts/ad_report/tests -v
  （build_deck を使うテストは python-pptx が無い環境では skip される）
"""
from __future__ import annotations

import csv
import json
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import unittest
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
AD = HERE.parent                         # 03_scripts/ad_report
REAL_VAULT = AD.parent.parent            # consulting-vault
SAMPLE = REAL_VAULT / "99_samples" / "ads" / "サンプル商事"
ANALYZE = AD / "ad_analyze.py"
BUILD = AD / "build_deck.py"
CHECK = AD / "check_deck.py"
MAKE_SAMPLE = AD / "tools" / "make_sample.py"
CLIENT = "サンプル商事"
MONTH = "2026-09"

sys.path.insert(0, str(AD))
import ad_common as C  # noqa: E402
import ad_analyze as A  # noqa: E402

try:
    import pptx  # noqa: F401
    from pptx import Presentation
    from pptx.oxml.ns import qn
    HAS_PPTX = True
except ImportError:  # pragma: no cover
    HAS_PPTX = False

PROFILE = """---
client: サンプル商事
aliases: [サンプル商事, サンプル]
slack_channels: [sample-shoji]
email_domains: [sample-shoji.example]
line_names: [サンプル商事 佐藤]
contract: 月額顧問
monthly_fee:
contact: 佐藤 様
style: です・ます調。
{targets}---

# サンプル商事 プロフィール（テスト用・架空）
"""


def run(script: Path, *args: str, python: str = sys.executable) -> subprocess.CompletedProcess:
    return subprocess.run([python, str(script)] + list(args), capture_output=True, text=True, cwd=str(REAL_VAULT))


def num(s: str) -> float:
    s = (s or "").strip().replace(",", "")
    return float(s) if s else 0.0


class VaultCase(unittest.TestCase):
    """一時 vault（クライアントフォルダ名はわざと NFD にする）を用意する。"""

    targets = "ad_kpi_targets: [cpa=4500, roas=1.5, ctr=1.0%]\n"
    minutes = True

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="adreport_test_"))
        self.vault = self.tmp / "vault"
        self.cdir = self.vault / "01_clients" / unicodedata.normalize("NFD", CLIENT)
        self.cdir.mkdir(parents=True)
        (self.cdir / "_profile.md").write_text(PROFILE.format(targets=self.targets), encoding="utf-8")
        if self.minutes:
            md = self.cdir / MONTH / "minutes"
            md.mkdir(parents=True)
            (md / "2026-09-17_月次定例.md").write_text("---\ntype: minutes\n---\n# 議事録（テスト）\n", encoding="utf-8")
        self.ads = self.cdir / MONTH / "ads"

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def analyze(self, data_dir: Path = SAMPLE / MONTH, *extra: str, python: str = sys.executable
                ) -> subprocess.CompletedProcess:
        return run(ANALYZE, "--client", CLIENT, "--month", MONTH, "--data-dir", str(data_dir),
                   "--vault", str(self.vault), *extra, python=python)

    def analysis(self) -> dict:
        return json.loads((self.ads / "analysis.json").read_text(encoding="utf-8"))

    def build(self, *extra: str) -> subprocess.CompletedProcess:
        return run(BUILD, "--client", CLIENT, "--month", MONTH, "--vault", str(self.vault), *extra)

    def check(self, *extra: str) -> subprocess.CompletedProcess:
        return run(CHECK, "--client", CLIENT, "--month", MONTH, "--vault", str(self.vault), *extra)

    def pptx_path(self) -> Path:
        return self.ads / f"{CLIENT}_{MONTH}_広告運用レポート.pptx"

    def write_insights(self, slides: dict) -> None:
        tpl = json.loads((self.ads / "insights.template.json").read_text(encoding="utf-8"))
        for k, v in slides.items():
            tpl["slides"][k]["text"] = v[0]
            tpl["slides"][k]["sources"] = v[1]
        (self.ads / "insights.json").write_text(json.dumps(tpl, ensure_ascii=False), encoding="utf-8")


# ---------------------------------------------------------------- サンプルデータ

class TestSample(unittest.TestCase):
    def test_sample_is_reproducible(self):
        """make_sample.py は乱数シード固定。作り直すと 99_samples と同じバイト列になる。"""
        tmp = Path(tempfile.mkdtemp(prefix="adsample_"))
        try:
            r = run(MAKE_SAMPLE, "--out", str(tmp))
            self.assertEqual(r.returncode, 0, r.stderr)
            for month in ("2026-08", "2026-09"):
                names = sorted(p.name for p in (SAMPLE / month).iterdir() if p.suffix == ".csv")
                self.assertEqual(len(names), 6)
                for n in names:
                    self.assertEqual((tmp / month / n).read_bytes(), (SAMPLE / month / n).read_bytes(), n)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_sample_shape(self):
        """キャンペーン3 × 広告セット3 × 広告3 = 27本、日別、日本語UIの列名・BOM付き。"""
        p = SAMPLE / MONTH / "広告_日別.csv"
        raw = p.read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.DictReader(raw.decode("utf-8-sig").splitlines()))
        ads = {(r["キャンペーン名"], r["広告セット名"], r["広告名"]) for r in rows}
        self.assertEqual(len(ads), 27)
        self.assertEqual(len({r["キャンペーン名"] for r in rows}), 3)
        self.assertIn("結果インジケーター", rows[0])
        self.assertEqual({r["結果インジケーター"] for r in rows if r["キャンペーン名"] == "認知_動画"}, {"reach"})


# ---------------------------------------------------------------- 集計

class TestAnalyze(VaultCase):
    def test_totals_match_independent_recalc(self):
        r = self.analyze()
        self.assertEqual(r.returncode, 0, r.stderr)
        a = self.analysis()
        with open(SAMPLE / MONTH / "広告_日別.csv", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        spend = sum(num(x["消化金額 (JPY)"]) for x in rows)
        imps = sum(num(x["インプレッション"]) for x in rows)
        link = sum(num(x["リンクのクリック"]) for x in rows)
        pur = sum(num(x["購入"]) for x in rows)
        pval = sum(num(x["購入のコンバージョン値"]) for x in rows)
        t = a["totals"]
        self.assertAlmostEqual(t["spend"]["value"], spend, places=2)
        self.assertEqual(t["impressions"]["value"], imps)
        self.assertEqual(t["link_clicks"]["value"], link)
        self.assertEqual(t["cv"]["value"], pur)
        self.assertAlmostEqual(t["purchase_value"]["value"], pval, places=2)
        self.assertAlmostEqual(t["cpa"]["value"], spend / pur, places=1)
        self.assertAlmostEqual(t["roas"]["value"], pval / spend, places=3)
        self.assertAlmostEqual(t["ctr"]["value"], link / imps, places=5)
        self.assertAlmostEqual(t["cpm"]["value"], spend / imps * 1000, places=1)
        self.assertEqual(t["spend"]["display"], f"¥{spend:,.0f}")
        # キャンペーン・広告の合計も一致
        self.assertAlmostEqual(sum(c["metrics"]["spend"]["value"] for c in a["campaigns"]), spend, places=2)
        self.assertAlmostEqual(sum(x["metrics"]["spend"]["value"] for x in a["ads"]), spend, places=2)
        self.assertEqual(len(a["ads"]), 27)
        # 出どころ（どのファイルの何行目か）
        self.assertEqual(a["provenance"]["total"][0]["row_count"], len(rows))

    def test_results_and_reach_rules(self):
        """目的の違う「結果」は合算しない。リーチは期間合計があれば実数、無ければ参考値。"""
        self.analyze()
        a = self.analysis()
        self.assertIsNone(a["totals"]["results"]["value"])            # 購入とリーチが混在
        self.assertTrue(a["totals"]["reach"].get("reference"))        # 日別合算は参考値
        c3 = next(c for c in a["campaigns"] if c["name"] == "認知_動画")
        self.assertFalse(c3["metrics"]["reach"].get("reference"))     # 期間合計_キャンペーン.csv の値
        with open(SAMPLE / MONTH / "期間合計_キャンペーン.csv", encoding="utf-8-sig", newline="") as f:
            per = {r["キャンペーン名"]: r for r in csv.DictReader(f)}
        self.assertEqual(c3["metrics"]["reach"]["value"], num(per["認知_動画"]["リーチ"]))
        self.assertEqual(c3["metrics"]["results"]["value"], num(per["認知_動画"]["結果"]))
        self.assertEqual(c3["result_label"], "リーチ")
        # 置換時は「（参考値）」が付く
        out, _ = C.resolve_placeholders("{{total.reach}}", a["values"])
        self.assertTrue(out.endswith("（参考値）"))

    def test_month_over_month(self):
        self.analyze()
        a = self.analysis()
        self.assertTrue(a["has_prev"])
        v = a["values"]
        diff = v["total.spend"]["value"] - v["prev.spend"]["value"]
        self.assertAlmostEqual(v["mom.spend.diff"]["value"], diff, places=2)
        self.assertAlmostEqual(v["mom.spend.pct"]["value"], diff / v["prev.spend"]["value"], places=5)
        self.assertEqual(v["mom.cpa.trend"]["value"], "改善" if v["mom.cpa.pct"]["value"] < 0 else "悪化")
        # 9月だけの新作は「新規」
        new = [x for x in a["ads"] if x["new"]]
        self.assertEqual([x["name"] for x in new], ["新規_UGC動画_9月新作"])

    def test_no_prev_month(self):
        tmp = self.tmp / "only"
        shutil.copytree(SAMPLE / MONTH, tmp / MONTH)
        r = self.analyze(tmp / MONTH)
        self.assertEqual(r.returncode, 0, r.stderr)
        a = self.analysis()
        self.assertFalse(a["has_prev"])
        self.assertNotIn("mom.spend.pct", a["values"])

    def test_targets(self):
        self.analyze()
        v = self.analysis()["values"]
        self.assertEqual(v["target.cpa.goal"]["value"], 4500)
        self.assertEqual(v["target.ctr.goal"]["value"], 0.01)      # 1.0% → 0.01
        cpa = v["total.cpa"]["value"]
        self.assertEqual(v["target.cpa.status"]["value"], "達成" if cpa <= 4500 else "未達")
        self.assertAlmostEqual(v["target.cpa.achievement"]["value"], 4500 / cpa, places=4)

    def test_zero_division_is_null(self):
        d = self.tmp / "zero" / MONTH
        d.mkdir(parents=True)
        with open(d / "ads.csv", "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["キャンペーン名", "広告セット名", "広告名", "日", "消化金額 (JPY)", "インプレッション",
                        "リンクのクリック", "購入", "購入のコンバージョン値"])
            w.writerow(["C", "S", "A", "2026-09-01", "0", "0", "0", "", ""])
            w.writerow(["C", "S", "B", "2026-09-02", "1000", "500", "0", "", ""])
        r = self.analyze(d)
        self.assertEqual(r.returncode, 0, r.stderr)
        a = self.analysis()
        v = a["values"]
        self.assertIsNone(v["day.2026-09-01.ctr"]["value"])
        self.assertEqual(v["day.2026-09-01.ctr"]["display"], "—")
        self.assertIsNone(v["day.2026-09-01.cpm"]["value"])
        self.assertIsNone(v["total.cpa"]["value"])                 # 購入 0 件
        self.assertIsNone(v["total.cpc"]["value"])                 # リンククリック 0
        self.assertEqual(v["total.cpa"]["display"], "—")
        names = {x["id"]: x["name"] for x in a["ads"]}
        self.assertEqual([names[i] for i in a["rankings"]["zero_cv"]], ["B"])   # 消化0の A は入らない
        self.assertNotIn("mom.spend.pct", v)

    def test_column_name_variants(self):
        m = A.ColumnMapper(C.load_json(AD / "column_map.json"))
        headers = ["Campaign name", "Ad Set Name", "Ad name", "Day", "Amount spent (JPY)", "Impressions",
                   "Link clicks", "Purchases", "Purchases conversion value", "Reporting starts",
                   "消化金額（ＪＰＹ）", " インプレッション ", "CTR (リンククリックスルー率)", "CPM(1,000インプレッション単価) (JPY)",
                   "クリック（全て）", "結果インジケーター", "アトリビューション設定", "謎の列"]
        hm = m.map_headers(headers)
        idx = hm["index"]
        for k in ("campaign_name", "adset_name", "ad_name", "date", "spend", "impressions", "link_clicks",
                  "purchases", "purchase_value", "date_start", "clicks_all", "result_indicator",
                  "attribution_setting"):
            self.assertIn(k, idx, k)
        self.assertEqual(hm["currency"], "JPY")
        self.assertIn("消化金額（ＪＰＹ）", hm["duplicates"])        # 英語の列と同じ標準キー → 2本目は重複扱い
        self.assertIn("CTR (リンククリックスルー率)", hm["ignored"])
        self.assertIn("CPM(1,000インプレッション単価) (JPY)", hm["ignored"])
        self.assertEqual(hm["unknown"], ["謎の列"])
        self.assertEqual(A.norm_col("消化金額（ＪＰＹ）"), A.norm_col("消化金額 (JPY)"))

    def test_english_export_end_to_end(self):
        d = self.tmp / "en" / MONTH
        d.mkdir(parents=True)
        with open(d / "export.csv", "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(["Campaign name", "Ad set name", "Ad name", "Day", "Amount spent (JPY)", "Impressions",
                        "Link clicks", "Purchases", "Purchases conversion value", "Unknown metric"])
            w.writerow(["Camp", "Set", "Ad1", "2026/9/1", "1,200", "1000", "20", "2", "5000", "x"])
            w.writerow(["Camp", "Set", "Ad1", "2026/9/2", "800", "900", "10", "", "", "y"])
        r = self.analyze(d)
        self.assertEqual(r.returncode, 0, r.stderr)
        a = self.analysis()
        self.assertEqual(a["totals"]["spend"]["value"], 2000)
        self.assertEqual(a["totals"]["cpa"]["value"], 1000)
        self.assertEqual(a["input"]["files"][0]["unknown_columns"], ["Unknown metric"])
        self.assertIn("未知の列", r.stderr)

    def test_xlsx_without_openpyxl(self):
        try:
            import openpyxl  # noqa: F401
            self.skipTest("openpyxl が入っている環境では確認できない")
        except ImportError:
            pass
        d = self.tmp / "x" / MONTH
        d.mkdir(parents=True)
        (d / "export.xlsx").write_bytes(b"PK\x03\x04dummy")
        r = self.analyze(d)
        self.assertEqual(r.returncode, 2)
        self.assertIn("openpyxl", r.stderr)
        self.assertIn("CSV", r.stderr)

    def test_template_and_next_actions_sources(self):
        self.analyze()
        tpl = json.loads((self.ads / "insights.template.json").read_text(encoding="utf-8"))
        slides = tpl["slides"]
        for sid in ("summary", "campaign.C1", "ad.C1S1A1", "age_gender", "issues", "next_actions"):
            self.assertIn(sid, slides)
            self.assertEqual(slides[sid]["text"], "")
        self.assertEqual(len([k for k in slides if k.startswith("ad.")]), 27)
        self.assertIn("total.cpa", slides["summary"]["placeholders"])
        self.assertEqual(slides["next_actions"]["sources"], [f"{MONTH}/minutes/2026-09-17_月次定例.md"])
        self.assertFalse((self.ads / "insights.json").exists())      # 雛形だけ。insights.json は作らない

    def test_runs_jsonl(self):
        self.analyze()
        lines = (self.vault / "04_logs" / "runs.jsonl").read_text(encoding="utf-8").strip().split("\n")
        rec = json.loads(lines[-1])
        self.assertEqual(rec["script"], "ad_analyze")
        self.assertEqual(rec["summary"]["client"], CLIENT)

    @unittest.skipUnless(Path("/usr/bin/python3").exists(), "macOS 標準の python3 が無い")
    def test_python39_stdlib_scripts(self):
        """ad_analyze.py / check_deck.py は macOS 標準の Python 3.9（標準ライブラリのみ）で動く。"""
        ver = subprocess.run(["/usr/bin/python3", "-c", "import sys;print(sys.version_info[:2])"],
                             capture_output=True, text=True).stdout
        if "(3, 9)" not in ver:
            self.skipTest(f"/usr/bin/python3 が 3.9 ではない: {ver}")
        r = self.analyze(python="/usr/bin/python3")
        self.assertEqual(r.returncode, 0, r.stderr)
        r2 = run(CHECK, "--client", CLIENT, "--month", MONTH, "--vault", str(self.vault), python="/usr/bin/python3")
        self.assertIn("pptx が無い", r2.stdout)   # pptx はまだ無いので ERROR（=動いている）


# ---------------------------------------------------------------- プレースホルダ・生の数字

class TestPlaceholders(unittest.TestCase):
    values = {"total.cpa": {"value": 3245.2, "unit": "JPY", "display": "¥3,245"},
              "mom.cpa.pct": {"value": -0.1, "unit": "ratio", "display": "−10.0%"}}

    def test_replace(self):
        out, unknown = C.resolve_placeholders("CPAは{{total.cpa}}（前月比{{ mom.cpa.pct }}）", self.values)
        self.assertEqual(out, "CPAは¥3,245（前月比−10.0%）")
        self.assertEqual(unknown, [])

    def test_unknown_key(self):
        _, unknown = C.resolve_placeholders("{{total.cpx}} と {{total.cpa}}", self.values)
        self.assertEqual(unknown, ["total.cpx"])
        _, unknown2 = C.resolve_placeholders("閉じていない {{total.cpa", self.values)
        self.assertTrue(unknown2)

    def test_raw_numbers(self):
        self.assertTrue(C.find_raw_numbers("CPAは3,000円で1.5%改善"))
        self.assertTrue(C.find_raw_numbers("CPAは￥３０００"))
        self.assertTrue(C.find_raw_numbers("ＣＰＡは２５００と"))
        self.assertTrue(C.find_raw_numbers("購入は１２件"))
        self.assertEqual(C.find_raw_numbers("CPAは{{total.cpa}}で、9月は好調（{{day.2026-09-21.label}}）"), [])


# ---------------------------------------------------------------- PPTX

@unittest.skipUnless(HAS_PPTX, "python-pptx が無い")
class TestBuild(VaultCase):
    def setUp(self) -> None:
        super().setUp()
        r = self.analyze()
        self.assertEqual(r.returncode, 0, r.stderr)

    def test_build_without_insights(self):
        r = self.build()
        self.assertEqual(r.returncode, 0, r.stderr)
        p = self.pptx_path()
        self.assertTrue(p.is_file())
        prs = Presentation(str(p))
        n = len(prs.slides)
        self.assertTrue(60 <= n <= 90, n)
        self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=2)
        texts = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
        self.assertIn(C.UNKNOWN_INSIGHT, texts)
        for i, s in enumerate(prs.slides, start=1):
            self.assertIn("【数字の出どころ】", s.notes_slide.notes_text_frame.text, f"slide {i}")
        c = self.check()
        self.assertEqual(c.returncode, 0, c.stdout + c.stderr)
        self.assertIn(f"スライド総数: {n}枚", c.stdout)

    def test_native_charts(self):
        self.build()
        prs = Presentation(str(self.pptx_path()))
        charts = [sh.chart for s in prs.slides for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart]
        self.assertGreater(len(charts), 50)
        kinds = {str(c.chart_type) for c in charts}
        self.assertTrue(any("COLUMN" in k for k in kinds))
        self.assertTrue(any("LINE" in k for k in kinds))
        self.assertTrue(any("BAR" in k for k in kinds))
        # グラフは単一種別（複合グラフを作っていない）
        for c in charts:
            self.assertEqual(len(c.plots), 1)

    def test_east_asian_font(self):
        self.build()
        font = json.loads((AD / "deck_config.json").read_text(encoding="utf-8"))["font"]["east_asian"]
        prs = Presentation(str(self.pptx_path()))
        n_rpr = 0
        for s in list(prs.slides)[:12]:
            for rpr in s._element.iter(qn("a:rPr")):
                ea = rpr.find(qn("a:ea"))
                self.assertIsNotNone(ea)
                self.assertEqual(ea.get("typeface"), font)
                n_rpr += 1
            for sh in s.shapes:
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    for d in sh.chart._chartSpace.iter(qn("a:defRPr")):
                        self.assertEqual(d.find(qn("a:ea")).get("typeface"), font)
        self.assertGreater(n_rpr, 100)
        with zipfile.ZipFile(self.pptx_path()) as z:
            theme = z.read("ppt/theme/theme1.xml").decode("utf-8")
        self.assertIn(f'<a:ea typeface="{font}"', theme)

    def test_placeholder_replaced_in_slide(self):
        v = self.analysis()["values"]
        self.write_insights({"summary": ("CPAは{{total.cpa}}、前月比{{mom.cpa.pct}}。", ["total.cpa", "mom.cpa"])})
        r = self.build()
        self.assertEqual(r.returncode, 0, r.stderr)
        prs = Presentation(str(self.pptx_path()))
        texts = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
        self.assertIn(f"CPAは{v['total.cpa']['display']}、前月比{v['mom.cpa.pct']['display']}。", texts)
        self.assertNotIn("{{", texts)
        self.assertEqual(self.check().returncode, 0)

    def test_unknown_placeholder_fails_build(self):
        self.write_insights({"summary": ("CPAは{{total.cpx}}", ["total.cpa"])})
        r = self.build()
        self.assertEqual(r.returncode, 1)
        self.assertIn("total.cpx", r.stderr)
        self.assertFalse(self.pptx_path().exists())
        c = self.check()
        self.assertEqual(c.returncode, 1)
        self.assertIn("解決できないプレースホルダ", c.stdout)

    def test_check_warns_raw_numbers_and_errors_on_missing_sources(self):
        self.write_insights({"summary": ("CPAは3,000円でした。", ["total.cpa"]),
                             "issues": ("{{total.cpa}} が課題", []),
                             "next_actions": ("議事録どおり", ["2026-09/minutes/存在しない.md"])})
        self.build()
        c = self.check()
        self.assertEqual(c.returncode, 1)
        self.assertIn("WARN", c.stdout)
        self.assertIn("3,000円", c.stdout)
        self.assertIn("sources が空", c.stdout)
        self.assertIn("存在しない", c.stdout)

    def test_template_option(self):
        tpl = self.tmp / "company.pptx"
        from pptx.util import Emu
        p = Presentation()
        p.slide_width, p.slide_height = Emu(12192000), Emu(6858000)
        s = p.slides.add_slide(p.slide_layouts[0])
        s.shapes.title.text = "テンプレの見本スライド"
        p.save(str(tpl))
        r = self.build("--template", str(tpl))
        self.assertEqual(r.returncode, 0, r.stderr)
        out = Presentation(str(self.pptx_path()))
        texts = "\n".join(sh.text_frame.text for s in out.slides for sh in s.shapes if sh.has_text_frame)
        self.assertNotIn("テンプレの見本スライド", texts)
        self.assertEqual(out.slide_width, 12192000)
        self.assertTrue(60 <= len(out.slides) <= 90)
        for s in out.slides:
            self.assertEqual(len(s.placeholders), 0)

    def test_config_sections_and_paging(self):
        self.build()
        base = len(Presentation(str(self.pptx_path())).slides)
        cfg = self.tmp / "cfg.json"
        conf = json.loads((AD / "deck_config.json").read_text(encoding="utf-8"))
        sections = [dict(s, enabled=(s["id"] != "ads")) for s in conf["sections"]]
        cfg.write_text(json.dumps({"sections": sections, "section_dividers": False}, ensure_ascii=False),
                       encoding="utf-8")
        r = self.build("--config", str(cfg))
        self.assertEqual(r.returncode, 0, r.stderr)
        n = len(Presentation(str(self.pptx_path())).slides)
        self.assertEqual(n, base - 27 - 7)   # 広告27枚と中扉7枚が消える
        cfg.write_text(json.dumps({"rows_per_page": {"daily_table": 8}}), encoding="utf-8")
        self.build("--config", str(cfg))
        prs = Presentation(str(self.pptx_path()))
        ids = [s.notes_slide.notes_text_frame.text.split("\n")[0] for s in prs.slides]
        self.assertEqual(len([i for i in ids if i.startswith("slide_id: daily.table")]), 4)   # 30日 ÷ 8行

    def test_next_actions_note_has_minutes_source(self):
        self.write_insights({"next_actions": ("議事録で決めた施策を進める。", [f"{MONTH}/minutes/2026-09-17_月次定例.md"])})
        self.build()
        prs = Presentation(str(self.pptx_path()))
        note = next(s.notes_slide.notes_text_frame.text for s in prs.slides
                    if s.notes_slide.notes_text_frame.text.startswith("slide_id: next_actions"))
        self.assertIn("2026-09-17_月次定例.md", note)
        self.assertEqual(self.check().returncode, 0)


if __name__ == "__main__":
    unittest.main()
