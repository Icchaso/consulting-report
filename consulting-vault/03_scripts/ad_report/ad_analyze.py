#!/usr/bin/env python3
"""Meta 広告のエクスポート（CSV/XLSX）を集計して analysis.json と所見の雛形を作るスクリプト。

大原則「ソースに遡れないことは書かない」の広告版:
  数字はすべてこのスクリプトが広告データから計算する。Claude は数字を書かず、所見の文章で
  {{campaign.C1.cpa}} のようなプレースホルダを使う（build_deck.py が analysis.json の値で置換する）。

何をするか:
  1. 入力フォルダの *.csv / *.xlsx を読み、列名を column_map.json で標準キーに寄せる
     （全角/半角・空白・通貨表記の揺れは正規化して照合。未知の列は警告して無視）
  2. 列の中身からファイルの種類を判定する（ファイル名は見ない）
       広告×日の明細（日別）/ 期間合計（キャンペーン等・日別でない）/
       内訳（年齢×性別・配置・デバイス・地域）
  3. 合計・日別・週別・曜日別・キャンペーン別・広告セット別・広告別・各内訳別・前月比・目標対比・
     ランキングを計算し、全部の値に単位と表示用文字列（例 "¥3,245" "1.82%"）を付ける
  4. <client>/<YYYY-MM>/ads/analysis.json と insights.template.json（所見の雛形）を書き出す

指標の定義（比率はすべて元の数値から再計算する。エクスポートの CTR/CPC/CPM/ROAS 列は使わない）:
  CTR（リンク）= リンククリック ÷ インプレッション   ※CTR（すべて）は クリック（すべて）÷ インプレッション（別物）
  CPC = 消化金額 ÷ リンククリック / CPM = 消化金額 ÷ インプレッション × 1000
  CV  = 購入（「購入」列がある場合。無ければ「結果」列。ただし結果の種類が違うものは合算しない）
  CVR = CV ÷ リンククリック / CPA = 消化金額 ÷ CV / ROAS = 購入のコンバージョン値 ÷ 消化金額
  リーチは足し算できない。期間合計のエクスポートがあればその値を使い、無ければ
  日別の合算を「参考値（日別合算・重複あり）」として出す（フリークエンシーも同様）。
  ゼロ除算は null（表示は「—」）。

使い方（vault ルートで実行）:
  python3 03_scripts/ad_report/ad_analyze.py --client サンプル商事 --month 2026-09
  python3 03_scripts/ad_report/ad_analyze.py --client サンプル商事 --month 2026-09 \\
      --data-dir 99_samples/ads/サンプル商事/2026-09 --vault /tmp/test-vault

  入力の既定は 01_clients/<client>/ads/<YYYY-MM>/。前月のフォルダ（…/ads/<前月>/、または
  --data-dir の兄弟フォルダ <前月>/）があれば前月比も計算する。--prev-data-dir で明示もできる。
  目標は _profile.md の frontmatter に `ad_kpi_targets: [cpa=3000, roas=3.0]` があれば読む。

標準ライブラリのみ・Python 3.9 で動作（xlsx は openpyxl がある場合だけ読む）。
実行ごとに 04_logs/runs.jsonl に1行追記する。
"""
from __future__ import annotations

import argparse
import calendar
import csv
import datetime as dt
import io
import json
import re
import sys
import unicodedata
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ad_common as C  # noqa: E402

SCRIPT_NAME = "ad_analyze"
SCHEMA_VERSION = 1

BASE_METRICS = ["spend", "impressions", "reach", "clicks_all", "link_clicks", "landing_page_views",
                "results", "purchases", "purchase_value", "adds_to_cart", "video_3s_views", "thruplays"]
TEXT_FIELDS = ["campaign_name", "adset_name", "ad_name", "date", "date_start", "date_stop", "result_indicator",
               "age", "gender", "platform", "placement", "device", "region", "country",
               "attribution_setting", "objective", "currency"]

# 指標の定義。unit: currency / count / ratio（0.0182 = 1.82%）/ x（倍）/ freq（回）
# better: up=大きいほど良い / down=小さいほど良い / neutral=良し悪しなし
METRIC_DEFS: Dict[str, dict] = {
    "spend": {"label": "消化金額", "unit": "currency", "better": "neutral", "formula": "エクスポートの消化金額の合計"},
    "impressions": {"label": "インプレッション", "unit": "count", "better": "neutral", "formula": "合計"},
    "reach": {"label": "リーチ", "unit": "count", "better": "neutral",
              "formula": "期間合計エクスポートの値。無い場合は日別の合算（重複を含む参考値）"},
    "frequency": {"label": "フリークエンシー", "unit": "freq", "better": "neutral", "formula": "インプレッション ÷ リーチ"},
    "clicks_all": {"label": "クリック（すべて）", "unit": "count", "better": "neutral",
                   "formula": "合計（いいね・プロフィールのクリック等も含む）"},
    "link_clicks": {"label": "リンククリック", "unit": "count", "better": "neutral", "formula": "合計"},
    "landing_page_views": {"label": "LPビュー", "unit": "count", "better": "neutral", "formula": "合計"},
    "results": {"label": "結果", "unit": "count", "better": "up",
                "formula": "合計（結果インジケーターが同じ範囲だけ。種類が混ざる範囲では算出しない）"},
    "purchases": {"label": "購入", "unit": "count", "better": "up", "formula": "合計"},
    "purchase_value": {"label": "購入金額", "unit": "currency", "better": "up", "formula": "購入のコンバージョン値の合計"},
    "adds_to_cart": {"label": "カート追加", "unit": "count", "better": "neutral", "formula": "合計"},
    "video_3s_views": {"label": "動画3秒再生", "unit": "count", "better": "neutral", "formula": "合計"},
    "thruplays": {"label": "ThruPlay", "unit": "count", "better": "neutral", "formula": "合計"},
    "cv": {"label": "CV", "unit": "count", "better": "up", "formula": "（cv_basis を参照）"},
    "ctr": {"label": "CTR（リンク）", "unit": "ratio", "better": "up", "formula": "リンククリック ÷ インプレッション"},
    "ctr_all": {"label": "CTR（すべて）", "unit": "ratio", "better": "up",
                "formula": "クリック（すべて）÷ インプレッション（リンクCTRとは別物）"},
    "cpc": {"label": "CPC（リンク）", "unit": "currency", "better": "down", "formula": "消化金額 ÷ リンククリック"},
    "cpm": {"label": "CPM", "unit": "currency", "better": "down", "formula": "消化金額 ÷ インプレッション × 1000"},
    "cvr": {"label": "CVR", "unit": "ratio", "better": "up", "formula": "CV ÷ リンククリック（分母はリンククリック）"},
    "cpa": {"label": "CPA", "unit": "currency", "better": "down", "formula": "消化金額 ÷ CV"},
    "roas": {"label": "ROAS", "unit": "x", "better": "up", "formula": "購入金額（コンバージョン値）÷ 消化金額"},
    "aov": {"label": "平均購入単価", "unit": "currency", "better": "up", "formula": "購入金額 ÷ 購入"},
    "cost_per_result": {"label": "結果の単価", "unit": "currency", "better": "down",
                        "formula": "消化金額 ÷ 結果（結果インジケーターが同じ範囲だけ）"},
}
CORE = ["spend", "impressions", "link_clicks", "ctr", "cpc", "cpm", "cv", "cvr", "cpa", "purchase_value", "roas"]
SHORT = ["spend", "cv", "cpa", "ctr", "cvr", "roas"]
MAIN_ORDER = ["spend", "impressions", "reach", "frequency", "clicks_all", "link_clicks", "ctr", "ctr_all", "cpc",
              "cpm", "landing_page_views", "cv", "cvr", "cpa", "purchases", "purchase_value", "roas", "aov",
              "results", "cost_per_result", "adds_to_cart", "video_3s_views", "thruplays"]

CURRENCIES = {"JPY": ("¥", 0), "USD": ("$", 2), "EUR": ("€", 2)}
CURRENCY_ALIASES = {"jpy": "JPY", "円": "JPY", "¥": "JPY", "usd": "USD", "$": "USD", "eur": "EUR", "€": "EUR"}
WEEKDAY_IDS = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]
WEEKDAY_JA = "月火水木金土日"
AGE_ORDER = ["13-17", "18-24", "25-34", "35-44", "45-54", "55-64", "65+"]
GENDER_ORDER = ["female", "male", "unknown"]
GENDER_IDS = {"女性": "female", "男性": "male", "不明": "unknown", "すべて": "all"}
BREAKDOWN_KINDS = ["age_gender", "placement", "device", "region"]
BREAKDOWN_LABELS = {"age_gender": "年齢×性別", "age": "年齢", "gender": "性別", "placement": "配置",
                    "device": "デバイス", "region": "地域"}
REACH_INDICATORS = ("reach",)


class InputError(Exception):
    """利用者が直すべき入力の問題（メッセージをそのまま表示する）。"""


# ---------------------------------------------------------------- 表示用の整形

class Fmt:
    def __init__(self, currency: str = "JPY") -> None:
        self.currency = currency
        self.symbol, self.decimals = CURRENCIES.get(currency, (currency + " ", 2))

    def unit_name(self, unit: str) -> str:
        return self.currency if unit == "currency" else unit

    def display(self, unit: str, v) -> str:
        if v is None:
            return C.NA_DISPLAY
        if unit == "currency":
            dec = 2 if (self.decimals == 0 and 0 < abs(v) < 10) else self.decimals  # ¥1.43 のような小さい単価
            return f"{self.symbol}{v:,.{dec}f}"
        if unit == "count":
            return f"{v:,.0f}"
        if unit == "ratio":
            return f"{v * 100:.2f}%"
        if unit == "x":
            return f"{v:.2f}倍"
        if unit == "freq":
            return f"{v:.2f}回"
        if unit == "pct":
            return f"{v * 100:.1f}%"
        return str(v)

    def diff_display(self, unit: str, d) -> str:
        if d is None:
            return C.NA_DISPLAY
        sign = "+" if d > 0 else ("−" if d < 0 else "±")
        a = abs(d)
        if unit == "currency":
            return f"{sign}{self.symbol}{a:,.{self.decimals}f}"
        if unit == "count":
            return f"{sign}{a:,.0f}"
        if unit == "ratio":
            return f"{sign}{a * 100:.2f}pt"
        if unit == "x":
            return f"{sign}{a:.2f}"
        if unit == "freq":
            return f"{sign}{a:.2f}"
        return f"{sign}{a}"

    def pct_change_display(self, p) -> str:
        if p is None:
            return C.NA_DISPLAY
        sign = "+" if p > 0 else ("−" if p < 0 else "±")
        return f"{sign}{abs(p) * 100:.1f}%"


def round_value(unit: str, v):
    if v is None:
        return None
    if unit == "count":
        r = round(v)
        return int(r) if abs(v - r) < 1e-9 else round(v, 4)
    if unit == "currency":
        return round(v, 2)
    if unit == "ratio":
        return round(v, 6)
    return round(v, 4)


def entry(fmt: Fmt, metric: str, v, reference: bool = False, note: Optional[str] = None) -> dict:
    unit = METRIC_DEFS[metric]["unit"]
    e = {"value": round_value(unit, v), "unit": fmt.unit_name(unit), "display": fmt.display(unit, v)}
    if reference:
        e["reference"] = True
    if note:
        e["note"] = note
    return e


def text_entry(s: str) -> dict:
    return {"value": s, "unit": "text", "display": s}


def plain_entry(fmt: Fmt, unit: str, v) -> dict:
    return {"value": round_value(unit, v), "unit": fmt.unit_name(unit), "display": fmt.display(unit, v)}


def subset(mt: Optional[Dict[str, dict]], keys: List[str]) -> Optional[Dict[str, dict]]:
    if mt is None:
        return None
    return {k: v for k, v in mt.items() if k in keys}


def safe_div(a, b):
    if a is None or b is None or b == 0:
        return None
    return a / b


# ---------------------------------------------------------------- 列名の照合

_CURRENCY_PAREN_RE = re.compile(r"[(\[]\s*(jpy|円|¥|usd|\$|eur|€)\s*[)\]]", re.I)


def norm_col(s: str) -> str:
    """列名の照合キー: NFKC → 小文字 → 通貨表記の除去 → 空白・括弧・区切り記号の除去。"""
    t = unicodedata.normalize("NFKC", str(s)).strip().lower()
    t = t.replace("\ufeff", "")
    t = _CURRENCY_PAREN_RE.sub("", t)
    t = re.sub(r"[\s_\-・:/,、.()\[\]]", "", t)
    t = t.replace("全て", "すべて")
    return t


def detect_currency(header: str) -> Optional[str]:
    t = unicodedata.normalize("NFKC", header).lower()
    m = _CURRENCY_PAREN_RE.search(t)
    if not m:
        return None
    return CURRENCY_ALIASES.get(m.group(1).lower())


class ColumnMapper:
    def __init__(self, cmap: dict) -> None:
        self.lookup: Dict[str, str] = {}
        for std, cands in cmap.get("columns", {}).items():
            if std.startswith("_"):
                continue
            for c in [std] + list(cands):
                self.lookup.setdefault(norm_col(c), std)
        self.ignore = {norm_col(c) for c in cmap.get("ignore", [])}
        self.values = {k: {norm_col(a): b for a, b in v.items()}
                       for k, v in cmap.get("values", {}).items() if not k.startswith("_") and isinstance(v, dict)}

    def map_headers(self, headers: List[str]) -> dict:
        mapping: Dict[str, int] = {}
        original: Dict[str, str] = {}
        unknown: List[str] = []
        ignored: List[str] = []
        duplicates: List[str] = []
        currency = None
        for i, h in enumerate(headers):
            h = str(h or "")
            if h.strip() == "":
                continue
            key = norm_col(h)
            std = self.lookup.get(key)
            if std is None:
                if key in self.ignore:
                    ignored.append(h)
                else:
                    unknown.append(h)
                continue
            if std in mapping:
                duplicates.append(h)
                continue
            mapping[std] = i
            original[std] = h
            if std in ("spend", "purchase_value"):
                currency = currency or detect_currency(h)
        return {"index": mapping, "original": original, "unknown": unknown, "ignored": ignored,
                "duplicates": duplicates, "currency": currency}

    def value_label(self, field: str, raw: str) -> str:
        raw = (raw or "").strip()
        if raw == "":
            return "不明"
        return self.values.get(field, {}).get(norm_col(raw), raw)


# ---------------------------------------------------------------- ファイルの読み込み

def parse_num(s) -> Optional[float]:
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    t = unicodedata.normalize("NFKC", str(s)).strip()
    if t in ("", "-", "—", "–", "n/a", "N/A", "null", "None"):
        return None
    t = t.replace(",", "").replace("¥", "").replace("$", "").replace("€", "").replace("円", "").replace("%", "")
    t = t.strip()
    try:
        return float(t)
    except ValueError:
        return None


_DATE_RES = [
    re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})"),
    re.compile(r"^(\d{4})/(\d{1,2})/(\d{1,2})"),
    re.compile(r"^(\d{4})年(\d{1,2})月(\d{1,2})日"),
    re.compile(r"^(\d{4})(\d{2})(\d{2})$"),
]


def parse_date(s) -> Optional[dt.date]:
    if s is None:
        return None
    if isinstance(s, dt.datetime):
        return s.date()
    if isinstance(s, dt.date):
        return s
    t = unicodedata.normalize("NFKC", str(s)).strip()
    for rx in _DATE_RES:
        m = rx.match(t)
        if m:
            try:
                return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            except ValueError:
                return None
    return None


def read_table(path: Path) -> Tuple[List[str], List[Tuple[int, List[str]]], str]:
    """戻り値: (見出し, [(行番号, 値の配列)], 読み方の説明)。行番号は元ファイルの行（見出し=1行目）。"""
    if path.suffix.lower() == ".xlsx":
        try:
            import openpyxl  # type: ignore
        except ImportError:
            raise InputError(
                f"{path.name}: xlsx を読むには openpyxl が必要です。広告マネージャで「CSV」形式で書き出し直すか、"
                "`pip3 install openpyxl` をしてから再実行してください。")
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        headers: List[str] = []
        data: List[Tuple[int, List[str]]] = []
        for i, row in enumerate(rows_iter, start=1):
            vals = ["" if v is None else (v.isoformat() if isinstance(v, (dt.date, dt.datetime)) else str(v))
                    for v in row]
            if i == 1:
                headers = vals
            elif any(v.strip() for v in vals):
                data.append((i, vals))
        wb.close()
        return headers, data, f"xlsx（シート「{ws.title}」）"
    raw = path.read_bytes()
    text = None
    used = ""
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        text, used = raw.decode("utf-16"), "utf-16"
    else:
        for enc in ("utf-8-sig", "cp932"):
            try:
                text, used = raw.decode(enc), enc
                break
            except UnicodeDecodeError:
                continue
    if text is None:
        raise InputError(f"{path.name}: 文字コードを判別できません（UTF-8 / Shift_JIS / UTF-16 のどれでもない）")
    text = text.lstrip("\ufeff")
    first_line = text.split("\n", 1)[0]
    delim = "\t" if first_line.count("\t") > first_line.count(",") else ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delim)
    headers = []
    data = []
    for row in reader:
        if not headers:
            headers = row
            continue
        if not any(c.strip() for c in row):
            continue
        data.append((reader.line_num, row))
    return headers, data, f"csv（{used}、区切り{'タブ' if delim == chr(9) else 'カンマ'}）"


def compress_rows(lines: Iterable[int]) -> str:
    nums = sorted(set(lines))
    if not nums:
        return ""
    parts = []
    start = prev = nums[0]
    for n in nums[1:]:
        if n == prev + 1:
            prev = n
            continue
        parts.append(f"{start}-{prev}" if start != prev else str(start))
        start = prev = n
    parts.append(f"{start}-{prev}" if start != prev else str(start))
    return ",".join(parts)


# ---------------------------------------------------------------- 集計の器

class Agg:
    """足し算できる値の合計と、どのファイルの何行目を足したか。"""

    def __init__(self) -> None:
        self.sums: Dict[str, float] = {k: 0.0 for k in BASE_METRICS}
        self.seen: set = set()
        self.refs: List[Tuple[int, int]] = []
        self.indicators: set = set()
        self.dates: set = set()
        self.n = 0

    def add(self, rec: dict) -> None:
        for k, v in rec["m"].items():
            self.sums[k] += v
            self.seen.add(k)
        self.refs.append(rec["src"])
        if rec.get("ind"):
            self.indicators.add(rec["ind"])
        if rec.get("d") is not None:
            self.dates.add(rec["d"])
        self.n += 1

    def merge(self, other: "Agg") -> None:
        for k in BASE_METRICS:
            self.sums[k] += other.sums[k]
        self.seen |= other.seen
        self.refs.extend(other.refs)
        self.indicators |= other.indicators
        self.dates |= other.dates
        self.n += other.n


class Ctx:
    """集計中に共有する情報（整形・CVの基準・警告・読んだファイル）。"""

    def __init__(self, fmt: Fmt, cv_basis: str, mapper: ColumnMapper) -> None:
        self.fmt = fmt
        self.cv_basis = cv_basis
        self.mapper = mapper
        self.files: List[dict] = []
        self.warnings: List[str] = []

    def indicator_label(self, ind: str) -> str:
        return self.mapper.value_label("result_indicator", ind) if ind else "不明"


def compute_metrics(ctx: Ctx, agg: Agg, reach_override: Optional[dict] = None) -> Dict[str, dict]:
    """Agg から指標（派生指標を含む）を計算して {指標: entry} を返す。無い列の指標は入れない。"""
    fmt = ctx.fmt
    s, seen = agg.sums, agg.seen
    out: Dict[str, dict] = {}

    def has(k: str) -> bool:
        return k in seen

    spend = s["spend"] if has("spend") else None
    imps = s["impressions"] if has("impressions") else None
    link = s["link_clicks"] if has("link_clicks") else None
    uniform = len(agg.indicators) <= 1
    ind = next(iter(agg.indicators)) if len(agg.indicators) == 1 else None
    reach_like = ind in REACH_INDICATORS if ind else False

    for k in ("spend", "impressions", "clicks_all", "link_clicks", "landing_page_views", "purchases",
              "purchase_value", "adds_to_cart", "video_3s_views", "thruplays"):
        if has(k):
            out[k] = entry(fmt, k, s[k])

    # リーチ・フリークエンシー
    if has("reach") or (reach_override and reach_override.get("reach") is not None):
        if reach_override and reach_override.get("reach") is not None:
            reach, ref, note = reach_override["reach"], False, None
        else:
            reach = s["reach"]
            ref = agg.n > 1
            note = "参考値（日別・内訳の合算。重複を含む）" if ref else None
        out["reach"] = entry(fmt, "reach", reach, reference=ref, note=note)
        if imps is not None:
            out["frequency"] = entry(fmt, "frequency", safe_div(imps, reach), reference=ref, note=note)

    # 結果（種類が同じ範囲だけ）
    results = None
    if has("results"):
        if uniform:
            results = s["results"]
            if reach_like and reach_override and reach_override.get("results") is not None:
                results = reach_override["results"]
                out["results"] = entry(fmt, "results", results)
            elif reach_like and agg.n > 1:
                out["results"] = entry(fmt, "results", results, reference=True,
                                       note="参考値（結果＝リーチの合算。重複を含む）")
            else:
                out["results"] = entry(fmt, "results", results)
            out["cost_per_result"] = entry(fmt, "cost_per_result", safe_div(spend, results))
        else:
            out["results"] = entry(fmt, "results", None, note="結果の種類が混在するため合算しない")
            out["cost_per_result"] = entry(fmt, "cost_per_result", None, note="結果の種類が混在")

    # CV
    cv = None
    if ctx.cv_basis == "purchases":
        cv = s["purchases"] if has("purchases") else None
    else:
        cv = results if (has("results") and uniform) else None
    if has(ctx.cv_basis) or ctx.cv_basis == "results":
        note = None if cv is not None else "結果の種類が混在するため算出しない"
        out["cv"] = entry(fmt, "cv", cv, note=note)
        out["cpa"] = entry(fmt, "cpa", safe_div(spend, cv))
        if link is not None:
            out["cvr"] = entry(fmt, "cvr", safe_div(cv, link))
    if imps is not None:
        if link is not None:
            out["ctr"] = entry(fmt, "ctr", safe_div(link, imps))
        if has("clicks_all"):
            out["ctr_all"] = entry(fmt, "ctr_all", safe_div(s["clicks_all"], imps))
        if spend is not None:
            out["cpm"] = entry(fmt, "cpm", None if not imps else spend / imps * 1000)
    if spend is not None and link is not None:
        out["cpc"] = entry(fmt, "cpc", safe_div(spend, link))
    if has("purchase_value") and spend is not None:
        out["roas"] = entry(fmt, "roas", safe_div(s["purchase_value"], spend))
        if has("purchases"):
            out["aov"] = entry(fmt, "aov", safe_div(s["purchase_value"], s["purchases"]))
    return {k: out[k] for k in MAIN_ORDER if k in out}


def compute_mom(ctx: Ctx, cur: Dict[str, dict], prev: Optional[Dict[str, dict]]) -> Optional[Dict[str, dict]]:
    if prev is None:
        return None
    fmt = ctx.fmt
    out: Dict[str, dict] = {}
    for k, e in cur.items():
        if k not in prev:
            continue
        unit = METRIC_DEFS[k]["unit"]
        a, b = e.get("value"), prev[k].get("value")
        diff = (a - b) if (a is not None and b is not None) else None
        pct = safe_div(diff, abs(b)) if (diff is not None and b not in (None, 0)) else None
        better = METRIC_DEFS[k]["better"]
        if pct is None:
            trend = C.NA_DISPLAY
        elif abs(pct) < 0.005:
            trend = "横ばい"
        elif better == "neutral":
            trend = "増加" if pct > 0 else "減少"
        elif (pct > 0) == (better == "up"):
            trend = "改善"
        else:
            trend = "悪化"
        out[k] = {
            "diff": {"value": round_value(unit, diff), "unit": fmt.unit_name(unit) if unit != "ratio" else "pt",
                     "display": fmt.diff_display(unit, diff)},
            "pct": {"value": None if pct is None else round(pct, 6), "unit": "ratio",
                    "display": fmt.pct_change_display(pct)},
            "trend": text_entry(trend),
        }
    return out


# ---------------------------------------------------------------- 入力の読み込みと分類

def load_inputs(ctx: Ctx, data_dir: Path, month: str, vault: Path) -> dict:
    """data_dir の全ファイルを読み、種類ごとのレコードに分ける。"""
    y, m = int(month[:4]), int(month[5:7])
    first = dt.date(y, m, 1)
    last = dt.date(y, m, calendar.monthrange(y, m)[1])
    out = {"daily": [], "period": [], "age_gender": [], "placement": [], "device": [], "region": []}
    paths = sorted([p for p in data_dir.iterdir()
                    if p.is_file() and p.suffix.lower() in (".csv", ".xlsx")
                    and not p.name.startswith((".", "~$"))], key=lambda p: C.nfc(p.name))
    currencies = set()
    for p in paths:
        headers, rows, how = read_table(p)
        hm = ctx.mapper.map_headers(headers)
        idx = hm["index"]
        fid = len(ctx.files)
        info = {"id": fid, "path": C.rel_to(p, vault), "name": C.nfc(p.name), "read_as": how,
                "kind": None, "level": None, "rows_total": len(rows), "rows_used": 0, "skipped": {},
                "columns": hm["original"], "ignored_columns": hm["ignored"],
                "unknown_columns": hm["unknown"], "duplicate_columns": hm["duplicates"],
                "currency": hm["currency"]}
        ctx.files.append(info)
        if hm["unknown"]:
            ctx.warnings.append(f"{info['name']}: 未知の列を無視しました（column_map.json に無い）: "
                                + " / ".join(hm["unknown"]))
        if hm["currency"]:
            currencies.add(hm["currency"])
        if "spend" not in idx or "impressions" not in idx:
            info["kind"] = "skipped"
            ctx.warnings.append(f"{info['name']}: 消化金額・インプレッションの列が見つからないため読み飛ばしました")
            continue

        def cell(row: List[str], key: str) -> str:
            i = idx.get(key)
            return row[i].strip() if i is not None and i < len(row) else ""

        # ファイルの種類
        dims = set(idx)
        if "age" in dims or "gender" in dims:
            kind = "age_gender"
        elif "placement" in dims or "platform" in dims:
            kind = "placement"
        elif "device" in dims:
            kind = "device"
        elif "region" in dims or "country" in dims:
            kind = "region"
        else:
            level = ("ad" if "ad_name" in dims else "adset" if "adset_name" in dims
                     else "campaign" if "campaign_name" in dims else "account")
            info["level"] = level
            is_daily = "date" in dims
            if not is_daily and "date_start" in dims:
                starts = {cell(r, "date_start") for _, r in rows}
                # 日別かどうか: 開始日と終了日が全行で同じで、行ごとに日付が違う
                is_daily = ("date_stop" in dims and len(starts) > 1
                            and all(cell(r, "date_start") == cell(r, "date_stop") for _, r in rows))
            if is_daily and level == "ad":
                kind = "daily"
            elif is_daily:
                kind = "skipped"
                ctx.warnings.append(f"{info['name']}: {level} 単位の日別データは未対応です（広告単位で日別に書き出してください）")
            else:
                kind = "period"
        info["kind"] = kind
        if kind == "skipped":
            continue

        seen_keys = set()
        for line_no, row in rows:
            def skip(reason: str) -> None:
                info["skipped"][reason] = info["skipped"].get(reason, 0) + 1

            names = [C.nfc(cell(row, k)) for k in ("campaign_name", "adset_name", "ad_name")]
            metrics: Dict[str, float] = {}
            for k in BASE_METRICS:
                if k in idx:
                    v = parse_num(cell(row, k))
                    metrics[k] = v if v is not None else 0.0
            rec = {"m": metrics, "src": (fid, line_no), "ind": cell(row, "result_indicator"),
                   "attr": cell(row, "attribution_setting"), "c": names[0] or "（名前なし）",
                   "s": names[1] or "（名前なし）", "a": names[2] or "（名前なし）", "d": None}
            if kind == "daily":
                if not any(names):
                    skip("キャンペーン・広告セット・広告名が空（合計行とみなす）")
                    continue
                d = parse_date(cell(row, "date")) if "date" in idx else None
                d = d or parse_date(cell(row, "date_start"))
                if d is None:
                    skip("日付が読めない")
                    continue
                if not (first <= d <= last):
                    skip("対象月の外")
                    continue
                key = (rec["c"], rec["s"], rec["a"], d)
                if key in seen_keys:
                    skip("同じ広告・同じ日の行が重複（2行目以降を無視）")
                    continue
                seen_keys.add(key)
                rec["d"] = d
                out["daily"].append(rec)
            else:
                ds = parse_date(cell(row, "date_start"))
                de = parse_date(cell(row, "date_stop"))
                if ds is not None and not (first <= ds <= last):
                    skip("対象月の外")
                    continue
                rec["period_ok"] = (ds in (None, first)) and (de in (None, last))
                rec["period"] = (ds.isoformat() if ds else None, de.isoformat() if de else None)
                if kind == "period":
                    if info["level"] != "account" and not any(names):
                        skip("名前が空（合計行とみなす）")
                        continue
                    rec["level"] = info["level"]
                    out["period"].append(rec)
                else:
                    if kind == "age_gender":
                        age = cell(row, "age") or "不明"
                        gender = ctx.mapper.value_label("gender", cell(row, "gender"))
                        rec["bucket"] = (age, gender)
                    elif kind == "placement":
                        plat = ctx.mapper.value_label("platform", cell(row, "platform")) if "platform" in idx else ""
                        plc = cell(row, "placement") if "placement" in idx else ""
                        rec["bucket"] = (plat, plc or "不明")
                    elif kind == "device":
                        rec["bucket"] = (ctx.mapper.value_label("device", cell(row, "device")),)
                    else:
                        rec["bucket"] = (cell(row, "region") or cell(row, "country") or "不明",)
                    out[kind].append(rec)
            info["rows_used"] += 1
        for reason, n in info["skipped"].items():
            ctx.warnings.append(f"{info['name']}: {n}行を読み飛ばしました（{reason}）")
    out["currencies"] = sorted(currencies)
    out["paths"] = paths
    return out


def detect_cv_basis(mapper: ColumnMapper, data_dir: Path) -> str:
    """「購入」列がある日別ファイルがあれば purchases、無ければ results。"""
    for p in sorted(data_dir.iterdir()):
        if p.suffix.lower() not in (".csv", ".xlsx") or p.name.startswith((".", "~$")):
            continue
        try:
            headers, _, _ = read_table(p)
        except InputError:
            continue
        hm = mapper.map_headers(headers)
        if "purchases" in hm["index"] and "ad_name" in hm["index"]:
            return "purchases"
    return "results"


# ---------------------------------------------------------------- 1か月分の集計

def month_bounds(month: str) -> Tuple[dt.date, dt.date]:
    y, m = int(month[:4]), int(month[5:7])
    return dt.date(y, m, 1), dt.date(y, m, calendar.monthrange(y, m)[1])


def day_label(d: dt.date) -> str:
    return f"{d.month}/{d.day}({WEEKDAY_JA[d.weekday()]})"


def slug(s: str) -> str:
    t = unicodedata.normalize("NFKC", s).lower().replace("+", "plus")
    t = re.sub(r"[^0-9a-z]+", "_", t).strip("_")
    return t


def analyze_month(ctx: Ctx, data: dict, month: str) -> dict:
    """1か月分のレコードを集計した中間結果（ID付け前）を返す。"""
    first, last = month_bounds(month)
    res: dict = {"total": Agg(), "campaigns": {}, "adsets": {}, "ads": {}, "days": {}, "ad_days": {},
                 "adset_days": {}, "campaign_days": {}, "period": {}, "breakdowns": {}, "attribution": {}}
    for rec in data["daily"]:
        c, s, a, d = rec["c"], rec["s"], rec["a"], rec["d"]
        res["total"].add(rec)
        res["campaigns"].setdefault(c, Agg()).add(rec)
        res["adsets"].setdefault((c, s), Agg()).add(rec)
        res["ads"].setdefault((c, s, a), Agg()).add(rec)
        res["days"].setdefault(d, Agg()).add(rec)
        res["campaign_days"].setdefault((c, d), Agg()).add(rec)
        res["adset_days"].setdefault((c, s, d), Agg()).add(rec)
        res["ad_days"].setdefault((c, s, a, d), Agg()).add(rec)
        if rec.get("attr"):
            res["attribution"].setdefault(c, set()).add(rec["attr"])
    # 期間合計（リーチの正しい値）
    for rec in data["period"]:
        if rec.get("attr"):
            res["attribution"].setdefault(rec["c"], set()).add(rec["attr"])
        if not rec.get("period_ok"):
            ctx.warnings.append(f"期間合計の行（{rec['c']}）の期間 {rec['period']} が対象月全体と一致しないため、リーチには使いません")
            continue
        level = rec["level"]
        key = {"account": ("__account__",), "campaign": (rec["c"],), "adset": (rec["c"], rec["s"]),
               "ad": (rec["c"], rec["s"], rec["a"])}[level]
        res["period"][(level,) + key] = rec
    for kind in BREAKDOWN_KINDS:
        recs = data[kind]
        if not recs:
            continue
        items: Dict[tuple, Agg] = {}
        by_camp: Dict[tuple, Dict[str, Agg]] = {}
        camp_totals: Dict[str, float] = {}
        for rec in recs:
            items.setdefault(rec["bucket"], Agg()).add(rec)
            by_camp.setdefault(rec["bucket"], {}).setdefault(rec["c"], Agg()).add(rec)
            camp_totals[rec["c"]] = camp_totals.get(rec["c"], 0.0) + rec["m"].get("spend", 0.0)
        res["breakdowns"][kind] = {"items": items, "by_campaign": by_camp, "campaign_spend": camp_totals}
    return res


def reach_override_for(res: dict, level: str, key: tuple) -> Optional[dict]:
    rec = res["period"].get((level,) + key)
    if rec is None or "reach" not in rec["m"]:
        return None
    ov = {"reach": rec["m"]["reach"], "src": rec["src"]}
    if rec.get("ind") in REACH_INDICATORS and "results" in rec["m"]:
        ov["results"] = rec["m"]["results"]
    return ov


def series(ctx: Ctx, days: List[dt.date], lookup) -> dict:
    """日別の配列（グラフ用の生の値）。"""
    out = {"dates": [d.isoformat() for d in days], "spend": [], "cv": [], "cpa": [], "ctr": [], "cvr": [],
           "roas": [], "impressions": [], "link_clicks": []}
    for d in days:
        agg = lookup(d)
        if agg is None:
            for k in out:
                if k != "dates":
                    out[k].append(0 if k in ("spend", "cv", "impressions", "link_clicks") else None)
            continue
        mt = compute_metrics(ctx, agg)
        for k in out:
            if k == "dates":
                continue
            e = mt.get(k)
            out[k].append(None if e is None else e["value"])
    return out


# ---------------------------------------------------------------- 目標

TARGET_ALIASES = {"cpa": "cpa", "roas": "roas", "ctr": "ctr", "cpc": "cpc", "cpm": "cpm", "cvr": "cvr",
                  "cv": "cv", "purchases": "cv", "購入": "cv", "spend": "spend", "予算": "spend",
                  "消化金額": "spend", "purchase_value": "purchase_value", "売上": "purchase_value",
                  "frequency": "frequency", "cost_per_result": "cost_per_result"}


def read_targets(client_dir: Path, vault: Path, warnings: List[str]) -> Tuple[Dict[str, float], Optional[str]]:
    p = client_dir / "_profile.md"
    if not p.is_file():
        return {}, None
    fm, _ = C.parse_frontmatter(p.read_text(encoding="utf-8"))
    raw = (fm or {}).get("ad_kpi_targets")
    if raw is None:
        return {}, C.rel_to(p, vault)
    items = raw if isinstance(raw, list) else [raw]
    targets: Dict[str, float] = {}
    for it in items:
        s = unicodedata.normalize("NFKC", str(it)).strip()
        if "=" not in s:
            warnings.append(f"_profile.md の ad_kpi_targets を読めません: {it}（metric=値 の形で書く）")
            continue
        k, _, v = s.partition("=")
        key = TARGET_ALIASES.get(k.strip().lower())
        if key is None:
            warnings.append(f"_profile.md の ad_kpi_targets の指標名が不明です: {k.strip()}")
            continue
        vs = v.strip()
        is_pct = vs.endswith("%")
        num = parse_num(vs)
        if num is None:
            warnings.append(f"_profile.md の ad_kpi_targets の値が数字ではありません: {it}")
            continue
        if METRIC_DEFS[key]["unit"] == "ratio" and (is_pct or num >= 1):
            num = num / 100.0
        targets[key] = num
    return targets, C.rel_to(p, vault)


# ---------------------------------------------------------------- 本体

def build_analysis(client: str, month: str, data_dir: Path, prev_dir: Optional[Path], client_dir: Path,
                   vault: Path, mapper: ColumnMapper, min_cv: int) -> Tuple[dict, List[str]]:
    cv_basis = detect_cv_basis(mapper, data_dir)
    probe = Ctx(Fmt("JPY"), cv_basis, mapper)
    data = load_inputs(probe, data_dir, month, vault)
    if not data["daily"]:
        raise InputError(
            f"{C.rel_to(data_dir, vault)} に「広告 × 日」の明細が見つかりません。広告マネージャで、"
            "レベル=広告・内訳=日（期間の分割=1日）で書き出した CSV を置いてください。"
            "（キャンペーン名・広告セット名・広告名・日・消化金額・インプレッションの列が必要）")
    currencies = data["currencies"]
    currency = currencies[0] if currencies else "JPY"
    ctx = Ctx(Fmt(currency), cv_basis, mapper)
    ctx.files, ctx.warnings = probe.files, probe.warnings
    if len(currencies) > 1:
        ctx.warnings.append(f"通貨が複数あります（{', '.join(currencies)}）。{currency} として表示します")
    fmt = ctx.fmt
    first, last = month_bounds(month)
    days = [first + dt.timedelta(days=i) for i in range((last - first).days + 1)]
    cur = analyze_month(ctx, data, month)

    prev = None
    prev_month = C.month_shift(month, -1)
    prev_ctx = None
    if prev_dir is not None and prev_dir.is_dir():
        prev_ctx = Ctx(Fmt(currency), cv_basis, mapper)
        try:
            pdata = load_inputs(prev_ctx, prev_dir, prev_month, vault)
            if pdata["daily"]:
                prev = analyze_month(prev_ctx, pdata, prev_month)
            else:
                ctx.warnings.append(f"前月フォルダ {C.rel_to(prev_dir, vault)} に広告×日の明細が無いため前月比は出しません")
        except InputError as e:
            ctx.warnings.append(f"前月データを読めませんでした: {e}")
        for w in prev_ctx.warnings:
            ctx.warnings.append("[前月] " + w)

    values: Dict[str, dict] = {}
    provenance: Dict[str, List[dict]] = {}
    file_paths = {f["id"]: f["path"] for f in ctx.files}
    prev_paths = {f["id"]: f["path"] for f in (prev_ctx.files if prev_ctx else [])}

    def prov(scope: str, agg: Agg, extra: Optional[List[Tuple[int, int]]] = None, prev_files: bool = False) -> None:
        paths = prev_paths if prev_files else file_paths
        by_file: Dict[int, List[int]] = {}
        for fid, line in list(agg.refs) + list(extra or []):
            by_file.setdefault(fid, []).append(line)
        provenance[scope] = [{"file": paths.get(fid, "?"), "rows": compress_rows(lines),
                              "row_count": len(set(lines))} for fid, lines in sorted(by_file.items())]

    def put_metrics(prefix: str, mt: Dict[str, dict], keys: Optional[List[str]] = None) -> None:
        for k, e in mt.items():
            if keys is None or k in keys:
                values[f"{prefix}.{k}"] = e

    def put_mom(prefix: str, mom: Optional[Dict[str, dict]], keys: Optional[List[str]] = None) -> None:
        if not mom:
            return
        for k, d in mom.items():
            if keys is None or k in keys:
                values[f"{prefix}.{k}.diff"] = d["diff"]
                values[f"{prefix}.{k}.pct"] = d["pct"]
                values[f"{prefix}.{k}.trend"] = d["trend"]

    # ---- 基本情報
    month_label = f"{first.year}年{first.month}月"
    period_label = f"{first.year}年{first.month}月{first.day}日〜{last.month}月{last.day}日"
    values["client"] = text_entry(client)
    values["month.label"] = text_entry(month_label)
    values["period.label"] = text_entry(period_label)
    values["period.days"] = {"value": len(days), "unit": "count", "display": f"{len(days)}日間"}
    values["prev_month.label"] = text_entry(f"{int(prev_month[:4])}年{int(prev_month[5:])}月")

    # ---- 全体
    acc_ov = reach_override_for(cur, "account", ("__account__",))
    total_m = compute_metrics(ctx, cur["total"], acc_ov)
    prev_total_m = compute_metrics(prev_ctx, prev["total"], reach_override_for(prev, "account", ("__account__",))) \
        if prev else None
    total_mom = compute_mom(ctx, total_m, prev_total_m)
    put_metrics("total", total_m)
    prov("total", cur["total"], [acc_ov["src"]] if acc_ov else None)
    if prev_total_m:
        put_metrics("prev", prev_total_m)
        prov("prev_total", prev["total"], prev_files=True)
        put_mom("mom", total_mom)
    ind_all = sorted(cur["total"].indicators)

    # ---- 日別・週別・曜日別
    daily_list = []
    for d in days:
        agg = cur["days"].get(d)
        mt = compute_metrics(ctx, agg) if agg else {}
        key = f"day.{d.isoformat()}"
        values[f"{key}.label"] = text_entry(day_label(d))
        put_metrics(key, mt, CORE)
        if agg:
            prov(key, agg)
        daily_list.append({"date": d.isoformat(), "label": day_label(d), "weekday": WEEKDAY_IDS[d.weekday()],
                           "has_data": agg is not None, "metrics": mt})
    total_series = series(ctx, days, lambda d: cur["days"].get(d))

    weeks: Dict[int, Agg] = {}
    week_days: Dict[int, List[dt.date]] = {}
    offset = first.weekday()
    for d in days:
        w = (d.day - 1 + offset) // 7 + 1
        week_days.setdefault(w, []).append(d)
        if d in cur["days"]:
            weeks.setdefault(w, Agg()).merge(cur["days"][d])
    weekly_list = []
    for w in sorted(week_days):
        ds = week_days[w]
        agg = weeks.get(w, Agg())
        mt = compute_metrics(ctx, agg) if agg.n else {}
        label = f"W{w} {ds[0].month}/{ds[0].day}〜{ds[-1].month}/{ds[-1].day}"
        key = f"week.W{w}"
        values[f"{key}.label"] = text_entry(label)
        values[f"{key}.days"] = {"value": len(ds), "unit": "count", "display": f"{len(ds)}日間"}
        put_metrics(key, mt, CORE)
        if agg.n:
            prov(key, agg)
        weekly_list.append({"id": f"W{w}", "label": label, "start": ds[0].isoformat(), "end": ds[-1].isoformat(),
                            "days": len(ds), "metrics": mt})
    weekday_list = []
    for wi, wid in enumerate(WEEKDAY_IDS):
        ds = [d for d in days if d.weekday() == wi]
        agg = Agg()
        for d in ds:
            if d in cur["days"]:
                agg.merge(cur["days"][d])
        mt = compute_metrics(ctx, agg) if agg.n else {}
        key = f"weekday.{wid}"
        values[f"{key}.label"] = text_entry(WEEKDAY_JA[wi] + "曜")
        values[f"{key}.days"] = {"value": len(ds), "unit": "count", "display": f"{len(ds)}日"}
        spend_pd = safe_div(agg.sums["spend"], len(ds)) if agg.n else None
        cv_val = mt.get("cv", {}).get("value")
        cv_pd = safe_div(cv_val, len(ds)) if cv_val is not None else None
        extra = {"spend_per_day": plain_entry(fmt, "currency", spend_pd),
                 "cv_per_day": {"value": None if cv_pd is None else round(cv_pd, 2), "unit": "count",
                                "display": C.NA_DISPLAY if cv_pd is None else f"{cv_pd:,.1f}"}}
        put_metrics(key, mt, CORE)
        values[f"{key}.spend_per_day"] = extra["spend_per_day"]
        values[f"{key}.cv_per_day"] = extra["cv_per_day"]
        if agg.n:
            prov(key, agg)
        weekday_list.append({"id": wid, "label": WEEKDAY_JA[wi], "days": len(ds), "metrics": mt, **extra})

    # ---- キャンペーン・広告セット・広告（ID は消化金額の多い順）
    def spend_of(agg: Agg) -> float:
        return agg.sums["spend"]

    campaigns_out, adsets_out, ads_out = [], [], []
    camp_names = sorted(cur["campaigns"], key=lambda c: (-spend_of(cur["campaigns"][c]), c))
    for ci, cname in enumerate(camp_names, start=1):
        cid = f"C{ci}"
        cagg = cur["campaigns"][cname]
        ov = reach_override_for(cur, "campaign", (cname,))
        mt = compute_metrics(ctx, cagg, ov)
        pm = None
        if prev and cname in prev["campaigns"]:
            pm = compute_metrics(prev_ctx, prev["campaigns"][cname],
                                 reach_override_for(prev, "campaign", (cname,)))
        mom = compute_mom(ctx, mt, pm)
        key = f"campaign.{cid}"
        ind = sorted(cagg.indicators)
        ind_label = " / ".join(ctx.indicator_label(i) for i in ind) if ind else "不明"
        values[f"{key}.name"] = text_entry(cname)
        values[f"{key}.result_label"] = text_entry(ind_label)
        put_metrics(key, mt)
        if pm:
            put_metrics(f"{key}.prev", pm, CORE)
            put_mom(f"{key}.mom", mom, CORE)
        prov(key, cagg, [ov["src"]] if ov else None)
        adset_names = sorted([k for k in cur["adsets"] if k[0] == cname],
                             key=lambda k: (-spend_of(cur["adsets"][k]), k[1]))
        adset_ids = []
        for si, skey in enumerate(adset_names, start=1):
            sid = f"{cid}S{si}"
            adset_ids.append(sid)
            sagg = cur["adsets"][skey]
            sov = reach_override_for(cur, "adset", skey)
            smt = compute_metrics(ctx, sagg, sov)
            spm = compute_metrics(prev_ctx, prev["adsets"][skey]) if prev and skey in prev["adsets"] else None
            smom = compute_mom(ctx, smt, spm)
            k2 = f"adset.{sid}"
            values[f"{k2}.name"] = text_entry(skey[1])
            values[f"{k2}.campaign_name"] = text_entry(cname)
            put_metrics(k2, smt)
            if spm:
                put_metrics(f"{k2}.prev", spm, CORE)
                put_mom(f"{k2}.mom", smom, CORE)
            prov(k2, sagg, [sov["src"]] if sov else None)
            ad_names = sorted([k for k in cur["ads"] if k[0] == cname and k[1] == skey[1]],
                              key=lambda k: (-spend_of(cur["ads"][k]), k[2]))
            ad_ids = []
            for ai, akey in enumerate(ad_names, start=1):
                aid = f"{sid}A{ai}"
                ad_ids.append(aid)
                aagg = cur["ads"][akey]
                aov = reach_override_for(cur, "ad", akey)
                amt = compute_metrics(ctx, aagg, aov)
                apm = compute_metrics(prev_ctx, prev["ads"][akey]) if prev and akey in prev["ads"] else None
                amom = compute_mom(ctx, amt, apm)
                k3 = f"ad.{aid}"
                values[f"{k3}.name"] = text_entry(akey[2])
                values[f"{k3}.adset_name"] = text_entry(skey[1])
                values[f"{k3}.campaign_name"] = text_entry(cname)
                active = sorted(aagg.dates)
                values[f"{k3}.active_days"] = {"value": len(active), "unit": "count", "display": f"{len(active)}日"}
                put_metrics(k3, amt)
                if apm:
                    put_metrics(f"{k3}.prev", apm, CORE)
                    put_mom(f"{k3}.mom", amom, CORE)
                prov(k3, aagg, [aov["src"]] if aov else None)
                ads_out.append({
                    "id": aid, "name": akey[2], "adset_id": sid, "campaign_id": cid,
                    "adset_name": skey[1], "campaign_name": cname, "metrics": amt,
                    "prev_metrics": subset(apm, CORE), "mom": subset(amom, CORE), "new": bool(prev) and apm is None,
                    "first_date": active[0].isoformat() if active else None,
                    "last_date": active[-1].isoformat() if active else None, "active_days": len(active),
                    "daily": series(ctx, days, lambda d, k=akey: cur["ad_days"].get(k + (d,))),
                })
            adsets_out.append({
                "id": sid, "name": skey[1], "campaign_id": cid, "campaign_name": cname, "metrics": smt,
                "prev_metrics": subset(spm, CORE), "mom": subset(smom, CORE), "new": bool(prev) and spm is None,
                "ad_ids": ad_ids,
                "daily": series(ctx, days, lambda d, k=skey: cur["adset_days"].get(k + (d,))),
            })
        attrs = sorted(cur["attribution"].get(cname, set()))
        campaigns_out.append({
            "id": cid, "name": cname, "metrics": mt, "prev_metrics": subset(pm, CORE), "mom": subset(mom, CORE),
            "new": bool(prev) and pm is None, "result_indicators": ind, "result_label": ind_label,
            "reach_source": "期間合計" if ov else "日別合算（参考値）", "attribution": attrs,
            "adset_ids": adset_ids,
            "daily": series(ctx, days, lambda d, k=cname: cur["campaign_days"].get((k, d))),
        })

    # 前月にあって今月に無いキャンペーン（停止）
    stopped = sorted(set(prev["campaigns"]) - set(cur["campaigns"])) if prev else []

    # ---- ランキング（広告）
    def mval(ad: dict, k: str):
        e = ad["metrics"].get(k)
        return None if e is None else e.get("value")

    with_cv = [a for a in ads_out if (mval(a, "cv") or 0) >= min_cv and mval(a, "cpa") is not None]
    rankings = {
        "cpa_best": [a["id"] for a in sorted(with_cv, key=lambda a: (mval(a, "cpa"), a["id"]))],
        "cpa_worst": [a["id"] for a in sorted(with_cv, key=lambda a: (-mval(a, "cpa"), a["id"]))],
        "spend": [a["id"] for a in sorted(ads_out, key=lambda a: (-(mval(a, "spend") or 0), a["id"]))],
        "roas_best": [a["id"] for a in sorted([a for a in ads_out if mval(a, "roas") is not None and
                                                (mval(a, "cv") or 0) >= min_cv],
                                               key=lambda a: (-mval(a, "roas"), a["id"]))],
        "zero_cv": [a["id"] for a in sorted([a for a in ads_out if (mval(a, "cv") or 0) == 0 and
                                              (mval(a, "spend") or 0) > 0],
                                             key=lambda a: (-(mval(a, "spend") or 0), a["id"]))],
    }
    ad_by_id = {a["id"]: a for a in ads_out}
    for kind, ids in rankings.items():
        for n, aid in enumerate(ids[:10], start=1):
            ad = ad_by_id[aid]
            k = f"rank.{kind}.{n}"
            values[f"{k}.name"] = text_entry(ad["name"])
            values[f"{k}.id"] = text_entry(aid)
            values[f"{k}.adset_name"] = text_entry(ad["adset_name"])
            put_metrics(k, ad["metrics"], SHORT)

    # ---- 内訳
    breakdowns_out: Dict[str, dict] = {}
    camp_id_by_name = {c["name"]: c["id"] for c in campaigns_out}
    for kind in BREAKDOWN_KINDS:
        b = cur["breakdowns"].get(kind)
        if not b:
            breakdowns_out[kind] = {"available": False}
            continue
        pb = prev["breakdowns"].get(kind) if prev else None
        # 広告×日の明細との整合（キャンペーン別の消化金額）
        for cname, sp in b["campaign_spend"].items():
            base = cur["campaigns"].get(cname)
            if base is None:
                ctx.warnings.append(f"内訳（{BREAKDOWN_LABELS[kind]}）に広告×日の明細に無いキャンペーンがあります: {cname}")
                continue
            ref = base.sums["spend"]
            if ref and abs(sp - ref) / ref > 0.01:
                ctx.warnings.append(f"内訳（{BREAKDOWN_LABELS[kind]}）の {cname} の消化金額が広告×日の明細と"
                                    f" {abs(sp - ref) / ref * 100:.1f}% ずれています（期間・フィルタの違いの可能性）")
        groups = {kind: b["items"]}
        prev_groups = {kind: pb["items"]} if pb else {}
        if kind == "age_gender":
            for sub, pos in (("age", 0), ("gender", 1)):
                g: Dict[tuple, Agg] = {}
                for bk, agg in b["items"].items():
                    g.setdefault((bk[pos],), Agg()).merge(agg)
                groups[sub] = g
                if pb:
                    pg: Dict[tuple, Agg] = {}
                    for bk, agg in pb["items"].items():
                        pg.setdefault((bk[pos],), Agg()).merge(agg)
                    prev_groups[sub] = pg
        total_spend = sum(a.sums["spend"] for a in b["items"].values())
        for gname, g in groups.items():
            items_out = []
            if gname in ("age_gender", "age", "gender"):
                def order_key(bk: tuple) -> tuple:
                    age_o = AGE_ORDER.index(bk[0]) if bk[0] in AGE_ORDER else len(AGE_ORDER)
                    if gname == "gender":
                        gid = GENDER_IDS.get(bk[0], bk[0])
                        return (GENDER_ORDER.index(gid) if gid in GENDER_ORDER else 9, bk[0])
                    if gname == "age":
                        return (age_o, bk[0])
                    gid = GENDER_IDS.get(bk[1], bk[1])
                    return (age_o, bk[0], GENDER_ORDER.index(gid) if gid in GENDER_ORDER else 9)
                keys = sorted(g, key=order_key)
            else:
                keys = sorted(g, key=lambda bk: (-g[bk].sums["spend"], bk))
            used_ids = set()
            for n, bk in enumerate(keys, start=1):
                if gname == "age_gender":
                    iid = f"a{slug(bk[0])}_{GENDER_IDS.get(bk[1], slug(bk[1]) or 'x')}"
                    label = f"{bk[0]} {bk[1]}"
                elif gname == "age":
                    iid, label = f"a{slug(bk[0])}", bk[0]
                elif gname == "gender":
                    iid, label = GENDER_IDS.get(bk[0], slug(bk[0]) or f"g{n}"), bk[0]
                else:
                    prefix = {"placement": "P", "device": "D", "region": "R"}[gname]
                    iid = f"{prefix}{n}"
                    label = " ".join(x for x in bk if x)
                if iid in used_ids or not iid:
                    iid = f"{iid or gname}_{n}"
                used_ids.add(iid)
                agg = g[bk]
                mt = compute_metrics(ctx, agg)
                share = safe_div(agg.sums["spend"], total_spend)
                pagg = prev_groups.get(gname, {}).get(bk) if prev_groups else None
                pm = compute_metrics(prev_ctx, pagg) if pagg else None
                mom = compute_mom(ctx, mt, pm)
                key = f"{gname}.{iid}"
                values[f"{key}.label"] = text_entry(label)
                values[f"{key}.spend_share"] = {"value": None if share is None else round(share, 6),
                                                "unit": "ratio", "display": fmt.display("pct", share)}
                put_metrics(key, mt)
                put_mom(f"{key}.mom", mom, SHORT)
                prov(key, agg)
                by_c = {}
                if gname == kind:
                    for cname, cagg in sorted(b["by_campaign"][bk].items()):
                        by_c[camp_id_by_name.get(cname, cname)] = subset(compute_metrics(ctx, cagg), SHORT)
                items_out.append({"id": iid, "label": label, "keys": list(bk), "metrics": mt,
                                  "spend_share": values[f"{key}.spend_share"], "prev_metrics": subset(pm, SHORT),
                                  "mom": subset(mom, SHORT),
                                  "by_campaign": by_c})
            breakdowns_out[gname] = {"available": True, "items": items_out,
                                     "files": sorted({provenance[f'{gname}.{i["id"]}'][0]["file"]
                                                      for i in items_out})}

    # ---- 目標
    warn_t: List[str] = []
    targets, profile_rel = read_targets(client_dir, vault, warn_t)
    ctx.warnings.extend(warn_t)
    target_items = []
    if targets:
        for k, goal in targets.items():
            actual = total_m.get(k, {}).get("value")
            better = METRIC_DEFS[k]["better"]
            if actual is None or goal == 0:
                ach, status = None, "判定不可"
            elif k == "spend":
                ach = actual / goal
                status = "予算内" if actual <= goal else "予算超過"
            elif better == "down":
                ach = goal / actual if actual else None
                status = "達成" if actual <= goal else "未達"
            else:
                ach = actual / goal
                status = "達成" if actual >= goal else "未達"
            key = f"target.{k}"
            values[f"{key}.goal"] = entry(fmt, k, goal)
            values[f"{key}.actual"] = total_m.get(k, entry(fmt, k, None))
            values[f"{key}.achievement"] = {"value": None if ach is None else round(ach, 6), "unit": "ratio",
                                            "display": fmt.display("pct", ach)}
            values[f"{key}.status"] = text_entry(status)
            provenance[key] = [{"file": profile_rel, "field": "ad_kpi_targets"}] + provenance.get("total", [])
            target_items.append({"metric": k, "label": METRIC_DEFS[k]["label"], "goal": values[f"{key}.goal"],
                                 "actual": values[f"{key}.actual"], "achievement": values[f"{key}.achievement"],
                                 "status": status, "better": better})
        values["target.status"] = text_entry("目標設定あり")
    else:
        values["target.status"] = text_entry("目標未設定")

    # ---- アトリビューション設定
    attr_all = sorted({a for s in cur["attribution"].values() for a in s})
    attribution = {
        "available": bool(attr_all), "values": attr_all,
        "by_campaign": {camp_id_by_name.get(c, c): sorted(v) for c, v in cur["attribution"].items()},
    }

    definitions = {k: {"label": v["label"], "unit": v["unit"], "better": v["better"], "formula": v["formula"]}
                   for k, v in METRIC_DEFS.items()}
    definitions["cv"]["formula"] = ("「購入」列の合計" if cv_basis == "purchases"
                                    else "「結果」列の合計（結果インジケーターが同じ範囲だけ）")
    definitions["cv"]["label"] = "CV（購入）" if cv_basis == "purchases" else "CV（結果）"

    analysis = {
        "schema_version": SCHEMA_VERSION,
        "client": client, "month": month, "prev_month": prev_month,
        "generated_at": C.now_iso(),
        "period": {"start": first.isoformat(), "end": last.isoformat(), "days": len(days), "label": period_label,
                   "month_label": month_label},
        "currency": currency,
        "cv_basis": cv_basis,
        "min_cv_for_ranking": min_cv,
        "definitions": definitions,
        "input": {"data_dir": C.rel_to(data_dir, vault),
                  "prev_data_dir": C.rel_to(prev_dir, vault) if (prev_dir and prev) else None,
                  "files": ctx.files, "prev_files": prev_ctx.files if prev_ctx else [],
                  "rows": {k: len(data[k]) for k in ["daily", "period"] + BREAKDOWN_KINDS},
                  "result_indicators": [{"raw": i, "label": ctx.indicator_label(i)} for i in ind_all],
                  "warnings": ctx.warnings},
        "has_prev": prev is not None,
        "targets": {"status": "set" if targets else "unset", "source": profile_rel, "items": target_items},
        "attribution": attribution,
        "totals": total_m, "prev_totals": prev_total_m, "mom": total_mom,
        "daily": daily_list, "daily_series": total_series,
        "weekly": weekly_list, "weekday": weekday_list,
        "campaigns": campaigns_out, "stopped_campaigns": stopped, "adsets": adsets_out, "ads": ads_out,
        "rankings": rankings,
        "breakdowns": breakdowns_out,
        "provenance": provenance,
    }
    analysis["insight_slots"] = build_slots(analysis, values, client_dir, month)
    analysis["values"] = values
    return analysis, ctx.warnings


def strip_slot_placeholders(analysis: dict) -> dict:
    """analysis.json には枠の一覧だけを残す（使えるキーの見本は insights.template.json に書く）。"""
    out = dict(analysis)
    out["insight_slots"] = [{k: v for k, v in s.items() if k != "placeholders"} for s in analysis["insight_slots"]]
    return out


# ---------------------------------------------------------------- 所見の枠（スライドID）

def _keys(values: Dict[str, dict], prefix: str, metrics: Optional[List[str]] = None,
          extra: Optional[List[str]] = None) -> List[str]:
    out = []
    for k in values:
        if not k.startswith(prefix + "."):
            continue
        rest = k[len(prefix) + 1:]
        if metrics is None or rest in metrics or (extra and rest in extra):
            out.append(k)
    return out


def build_slots(analysis: dict, values: Dict[str, dict], client_dir: Path, month: str) -> List[dict]:
    """所見を書ける枠（= insights.json のスライドID）と、そこで使うと良いプレースホルダ。"""
    main_keys = CORE + ["reach", "frequency", "results", "cost_per_result", "ctr_all", "aov", "purchases"]
    mom_core = [f"{m}.{x}" for m in CORE for x in ("pct", "trend")]
    slots = []

    def add(sid: str, title: str, section: str, ph: List[str], hint: str, max_chars: int,
            sources: Optional[List[str]] = None) -> None:
        slots.append({"id": sid, "title": title, "section": section, "placeholders": ph, "hint": hint,
                      "max_chars": max_chars, "sources": sources or []})

    common = ["period.label", "month.label", "prev_month.label"]
    add("summary", "エグゼクティブサマリー", "summary",
        common + _keys(values, "total", main_keys) + _keys(values, "mom", mom_core) + ["target.status"]
        + [k for k in values if k.startswith("target.")],
        "今月の結論を2〜4文で（何が良く何が課題か）。数字は必ずプレースホルダで。", 220)
    add("targets", "目標対比", "targets", ["target.status"] + [k for k in values if k.startswith("target.")],
        "目標に対する達成状況。目標未設定なら空欄のままでよい。", 180)
    add("mom", "前月比較", "mom", _keys(values, "total", CORE) + _keys(values, "prev", CORE)
        + _keys(values, "mom", [f"{m}.{x}" for m in CORE for x in ("diff", "pct", "trend")]),
        "前月から何が変わったか。前月データが無ければ空欄のまま。", 200)
    add("daily", "日別推移", "daily", [k for k in values if k.startswith("day.") and
                                    k.rsplit(".", 1)[-1] in ("label", "spend", "cv", "cpa", "roas")],
        "日別の山谷と、その日付。日付も {{day.YYYY-MM-DD.label}} で書ける。", 200)
    add("weekly", "週別推移", "weekly", [k for k in values if k.startswith("week.")],
        "週ごとの変化。", 160)
    add("weekday", "曜日別", "weekly", [k for k in values if k.startswith("weekday.")],
        "曜日による差（1日あたりの値で比べる）。", 160)
    add("campaigns", "キャンペーン比較", "campaign_compare",
        [k for k in values if k.startswith("campaign.") and k.split(".")[2] in ("name", "spend", "cv", "cpa",
                                                                             "roas", "ctr", "result_label")
         and len(k.split(".")) == 3],
        "キャンペーン間の役割と成果の違い。目的（結果の種類）が違うものはCPAで単純比較しない。", 200)
    for c in analysis["campaigns"]:
        cid = c["id"]
        ph = _keys(values, f"campaign.{cid}", main_keys, ["name", "result_label"]) \
            + _keys(values, f"campaign.{cid}.mom", mom_core)
        for sid in c["adset_ids"]:
            ph += _keys(values, f"adset.{sid}", ["name", "spend", "cv", "cpa", "roas"])
        add(f"campaign.{cid}", f"キャンペーン: {c['name']}", "campaigns", ph,
            "このキャンペーンの成果と、広告セット間の差。", 200)
    for s in analysis["adsets"]:
        sid = s["id"]
        ph = _keys(values, f"adset.{sid}", main_keys, ["name", "campaign_name"]) \
            + _keys(values, f"adset.{sid}.mom", mom_core)
        for aid in s["ad_ids"]:
            ph += _keys(values, f"ad.{aid}", ["name", "spend", "cv", "cpa", "ctr", "roas"])
        add(f"adset.{sid}", f"広告セット: {s['name']}（{s['campaign_name']}）", "adsets", ph,
            "この広告セットの成果と、広告間の差。", 160)
    for a in analysis["ads"]:
        aid = a["id"]
        ph = _keys(values, f"ad.{aid}", main_keys, ["name", "adset_name", "campaign_name", "active_days"]) \
            + _keys(values, f"ad.{aid}.mom", mom_core)
        add(f"ad.{aid}", f"広告: {a['name']}（{a['adset_name']}）", "ads", ph,
            "このクリエイティブの評価（続ける/改善/止める の判断材料）。", 130)
    add("ranking", "上位・下位クリエイティブ", "ranking", [k for k in values if k.startswith("rank.")],
        "CPAの良い/悪いクリエイティブの共通点。", 200)
    for kind in ("age_gender", "placement", "device", "region"):
        if not analysis["breakdowns"].get(kind, {}).get("available"):
            continue
        groups = [kind] + (["age", "gender"] if kind == "age_gender" else [])
        ph = [k for k in values if k.split(".")[0] in groups and
              k.rsplit(".", 1)[-1] in ("label", "spend", "spend_share", "cv", "cpa", "ctr", "cvr", "roas")
              and ".mom." not in k]
        add(kind, BREAKDOWN_LABELS[kind] + "別", kind, ph, "どの層・面・端末で成果が出ているか。", 200)
    add("issues", "所見・課題", "issues", common + _keys(values, "total", CORE),
        "今月の課題を箇条書き（改行で区切る）。どれも数字はプレースホルダで。", 480)
    rec_files = []
    for sub in ("minutes", "proposals", "qa"):
        d = client_dir / month / sub
        if d.is_dir():
            for p in sorted(d.iterdir(), key=lambda x: C.nfc(x.name)):
                if p.is_file() and p.suffix.lower() == ".md":
                    rec_files.append(f"{month}/{sub}/{C.nfc(p.name)}")
    add("next_actions", "次月施策", "next_actions", common,
        "議事録・提案書で実際に話した施策だけを書く（sources にそのファイルを入れる）。記録が無ければ空欄のまま。", 480,
        rec_files)
    return slots


def write_template(path: Path, analysis: dict) -> None:
    values = analysis["values"]
    defs = analysis["definitions"]

    def hint(k: str) -> str:
        e = values[k]
        metric = k.rsplit(".", 1)[-1]
        label = defs[metric]["label"] if metric in defs else ""
        ref = "・参考値" if e.get("reference") else ""
        return f"{e['display']}（{label}{ref}）" if label else e["display"]

    tpl = {
        "_説明": [
            "このファイルは ad_analyze.py が作る雛形。コピーして insights.json という名前で同じフォルダに置き、text と sources だけを書く。",
            "text に数字を直接書かない。数字は必ず {{キー}} で書く（例: CPAは{{total.cpa}}で、前月比{{mom.cpa.pct}}）。build_deck.py が analysis.json の表示用の値で置き換える。",
            "使えるキーは analysis.json の values の全キー。placeholders はそのスライドで使いそうなキーと現在の値の見本（どのスライドでも全キーが使える）。",
            "存在しないキーを書くとビルドがエラーで止まる。数字＋単位（例「3,000円」「1.5%」）を直接書くと check_deck.py が WARN を出す。",
            "sources には根拠を書く（必須）。analysis.json のキー（例 total.cpa / campaign.C1）か、クライアントフォルダからの相対パス（例 2026-09/minutes/2026-09-03_月次定例.md）。",
            "書けないスライドは text を空のままにする（スライドに「所見：要確認」と出る）。それらしい文章で埋めない。",
            "text は改行で段落・箇条書きになる。max_chars は目安の文字数（置換後）。",
        ],
        "schema_version": SCHEMA_VERSION,
        "client": analysis["client"], "month": analysis["month"],
        "analysis_generated_at": analysis["generated_at"],
        "slides": {},
    }
    for s in analysis["insight_slots"]:
        tpl["slides"][s["id"]] = {
            "title": s["title"], "text": "", "sources": list(s["sources"]), "hint": s["hint"],
            "max_chars": s["max_chars"], "placeholders": {k: hint(k) for k in s["placeholders"] if k in values},
        }
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        json.dump(tpl, f, ensure_ascii=False, indent=1)  # 1キー1行（Claude が読みやすいように）
        f.write("\n")


# ---------------------------------------------------------------- CLI

def load_column_map(path: Optional[str]) -> ColumnMapper:
    p = Path(path) if path else C.script_dir() / "column_map.json"
    return ColumnMapper(C.load_json(p))


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Meta 広告のエクスポートを集計して analysis.json と所見の雛形を作る")
    ap.add_argument("--client", required=True, help="クライアント名（01_clients/ のフォルダ名）")
    ap.add_argument("--month", help="対象月 YYYY-MM（省略時は前月）")
    ap.add_argument("--data-dir", help="入力フォルダ（既定: 01_clients/<client>/ads/<YYYY-MM>/）")
    ap.add_argument("--prev-data-dir", help="前月の入力フォルダ（既定: 入力フォルダの兄弟 <前月>/ か 01_clients/<client>/ads/<前月>/）")
    ap.add_argument("--vault", help="vault ルート（省略時はこのスクリプトから自動判定）")
    ap.add_argument("--column-map", help="列名マッピング（既定: 03_scripts/ad_report/column_map.json）")
    ap.add_argument("--min-cv", type=int, default=3, help="CPAランキングに入れる最低CV数（既定 3）")
    args = ap.parse_args(argv)

    month = args.month or C.previous_month_of_today()
    if not C.MONTH_RE.match(month):
        print(f"エラー: 月は YYYY-MM 形式で指定してください（受け取った値: {month}）", file=sys.stderr)
        return 2
    vault = C.resolve_vault(args.vault)
    client = C.nfc(args.client.strip())
    client_dir = C.find_client_dir(vault, client)
    if client_dir is None:
        print(f"エラー: クライアント「{client}」のフォルダが {vault / '01_clients'} にありません", file=sys.stderr)
        return 2
    client = C.nfc(client_dir.name)

    def resolve(p: str) -> Path:
        q = Path(p).expanduser()
        return q if q.is_absolute() else (Path.cwd() / q)

    data_dir = resolve(args.data_dir) if args.data_dir else client_dir / "ads" / month
    if not data_dir.is_dir():
        print(f"エラー: 入力フォルダがありません: {data_dir}\n"
              f"  → 広告マネージャのエクスポートを 01_clients/{client}/ads/{month}/ に置くか、--data-dir で指定してください",
              file=sys.stderr)
        return 2
    prev_month = C.month_shift(month, -1)
    if args.prev_data_dir:
        prev_dir: Optional[Path] = resolve(args.prev_data_dir)
    elif args.data_dir and C.nfc(data_dir.name) == month:
        prev_dir = data_dir.parent / prev_month
    else:
        prev_dir = client_dir / "ads" / prev_month
    if prev_dir is not None and not prev_dir.is_dir():
        prev_dir = None

    try:
        mapper = load_column_map(args.column_map)
        analysis, warnings = build_analysis(client, month, data_dir, prev_dir, client_dir, vault, mapper,
                                            args.min_cv)
    except InputError as e:
        print(f"エラー: {e}", file=sys.stderr)
        C.append_run_log(vault, SCRIPT_NAME, {"client": client, "month": month, "error": str(e)})
        return 2

    out_dir = client_dir / month / "ads"
    out_path = out_dir / "analysis.json"
    tpl_path = out_dir / "insights.template.json"
    write_template(tpl_path, analysis)
    C.write_json(out_path, strip_slot_placeholders(analysis))

    t = analysis["totals"]
    print(f"ad_analyze: {client} {month}")
    for f in analysis["input"]["files"]:
        print(f"  読込: {f['name']}（{f['kind']}{'・' + f['level'] if f['level'] else ''}、"
              f"{f['rows_used']}/{f['rows_total']}行）")
    if analysis["has_prev"]:
        print(f"  前月: {analysis['input']['prev_data_dir']}")
    else:
        print("  前月: なし（前月比は出しません）")
    print(f"  合計: 消化金額 {t['spend']['display']} / CV {t.get('cv', {}).get('display', '—')}"
          f" / CPA {t.get('cpa', {}).get('display', '—')} / ROAS {t.get('roas', {}).get('display', '—')}")
    print(f"  キャンペーン {len(analysis['campaigns'])} / 広告セット {len(analysis['adsets'])} / "
          f"広告 {len(analysis['ads'])} / 所見の枠 {len(analysis['insight_slots'])}")
    print(f"  目標: {'あり' if analysis['targets']['status'] == 'set' else '目標未設定'}")
    for w in warnings:
        print(f"  [WARN] {w}", file=sys.stderr)
    print(f"  → {C.rel_to(out_path, vault)}")
    print(f"  → {C.rel_to(tpl_path, vault)}（コピーして insights.json を書く）")
    C.append_run_log(vault, SCRIPT_NAME, {
        "client": client, "month": month, "data_dir": analysis["input"]["data_dir"],
        "files": len(analysis["input"]["files"]), "rows": analysis["input"]["rows"],
        "has_prev": analysis["has_prev"], "warnings": len(warnings), "path": C.rel_to(out_path, vault),
    })
    return 0


if __name__ == "__main__":
    sys.exit(main())
