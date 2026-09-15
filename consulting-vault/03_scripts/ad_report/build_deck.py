#!/usr/bin/env python3
"""analysis.json（数字）と insights.json（所見の文章）から広告運用レポートの PowerPoint を作るスクリプト。

- 数字はすべて analysis.json の値（ad_analyze.py が広告データから計算したもの）をそのまま使う。
- 所見は insights.json の text を使い、{{キー}} を analysis.json の表示用の値で置き換える。
  存在しないキーがあればエラーで止まる（pptx は書かない）。
- insights.json が無い・text が空のスライドは「所見：要確認」を目立つ形で出す（空欄を埋めない）。
- グラフは PowerPoint のネイティブグラフ（後で数値を編集できる）。棒＋折れ線の2軸複合は使わない。
- 各スライドのノートに「このスライドの数字の出どころ」（analysis.json のキーと元ファイルの行）を書く。
- 日本語フォントは東アジアフォント（a:ea）にも設定する。

使い方（vault ルートで実行）:
  python3 03_scripts/ad_report/build_deck.py --client サンプル商事 --month 2026-09
  python3 03_scripts/ad_report/build_deck.py --client サンプル商事 --month 2026-09 \\
      --template 会社テンプレート.pptx --config my_deck_config.json --vault /path/to/vault

  入力: 01_clients/<client>/<YYYY-MM>/ads/analysis.json と insights.json（無くてもよい）
  出力: 01_clients/<client>/<YYYY-MM>/ads/<client>_<YYYY-MM>_広告運用レポート.pptx

python-pptx が必要（pip3 install -r 03_scripts/ad_report/requirements.txt）。コードは Python 3.9 互換。
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import math
import sys
from pathlib import Path
from typing import Callable, Dict, List, Optional, Sequence, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ad_common as C  # noqa: E402

try:
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Emu, Pt
except ImportError:  # pragma: no cover
    print("エラー: python-pptx がありません。`pip3 install -r 03_scripts/ad_report/requirements.txt` を実行してください。",
          file=sys.stderr)
    sys.exit(3)

SCRIPT_NAME = "build_deck"
W_IN, H_IN = 13.333, 7.5          # 設計上のスライド寸法（16:9）。テンプレートの寸法に合わせて拡大縮小する
ML, MR = 0.5, 0.5
CW = W_IN - ML - MR
CT, CB = 1.32, 6.88               # 本文の上端・下端
CH = CB - CT
LEFT_W = 8.3                      # 右に所見を置くときの左側の幅
INS_X = ML + LEFT_W + 0.22
INS_W = W_IN - MR - INS_X
NOTES_MARK = "【数字の出どころ】"
INSIGHT_EMPTY_NOTE = "insights.json 未記入"


# ---------------------------------------------------------------- 設定

def deep_merge(base: dict, over: dict) -> dict:
    out = copy.deepcopy(base)
    for k, v in over.items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = deep_merge(out[k], v)
        else:
            out[k] = copy.deepcopy(v)
    return out


def load_config(path: Optional[str]) -> dict:
    cfg = C.load_json(C.script_dir() / "deck_config.json")
    if path:
        cfg = deep_merge(cfg, C.load_json(Path(path)))
    return cfg


# ---------------------------------------------------------------- 描画の道具

PCT_FMT = '0.0"%"'  # 比率はグラフ上では百分率の数値（1.82 = 1.82%）で持つ（編集しやすく、どのビューアでも読める）


def pct100(v):
    return None if v is None else round(v * 100, 4)


def nice_max(v: float) -> float:
    """軸の最大値を切りのいい値にする（例 112,345 → 120,000）。"""
    if v <= 0:
        return 1.0
    mag = 10 ** math.floor(math.log10(v))
    for m in (1, 1.2, 1.5, 2, 2.5, 3, 4, 5, 6, 8, 10):
        if v <= m * mag:
            return m * mag
    return 10 * mag


def set_plot_layout(chart, x: float, y: float, w: float, h: float) -> None:
    """グラフの描画領域（軸の内側）を割合で固定する。横棒グラフで項目名の置き場を確保するため。"""
    pa = chart._chartSpace.find(".//" + qn("c:plotArea"))
    if pa is None:
        return
    old = pa.find(qn("c:layout"))
    if old is not None:
        pa.remove(old)
    layout = OxmlElement("c:layout")
    ml = OxmlElement("c:manualLayout")
    for tag, val in (("c:layoutTarget", "inner"), ("c:xMode", "edge"), ("c:yMode", "edge"),
                     ("c:x", f"{x:.4f}"), ("c:y", f"{y:.4f}"), ("c:w", f"{w:.4f}"), ("c:h", f"{h:.4f}")):
        el = OxmlElement(tag)
        el.set("val", val)
        ml.append(el)
    layout.append(ml)
    pa.insert(0, layout)


def set_tick_skip(chart, n: int) -> None:
    """カテゴリ軸のラベルを n 個おきに表示する（日別グラフの日付の重なり防止）。"""
    cat = chart._chartSpace.find(".//" + qn("c:catAx"))
    if cat is None or n <= 1:
        return
    for tag in ("c:tickLblSkip", "c:tickMarkSkip"):
        old = cat.find(qn(tag))
        if old is not None:
            cat.remove(old)
    el = OxmlElement("c:tickLblSkip")
    el.set("val", str(n))
    for succ in ("c:noMultiLvlLbl", "c:extLst"):
        nxt = cat.find(qn(succ))
        if nxt is not None:
            nxt.addprevious(el)
            break
    else:
        cat.append(el)


def short_label(s: str, max_units: float = 15.0) -> str:
    """グラフの項目名を短くする（長い名前が1つあると全部の項目名が縮むビューアがあるため）。表では全文を出す。"""
    s = str(s)
    if C.text_units(s) <= max_units:
        return s
    out = ""
    for ch in s:
        if C.text_units(out + ch) > max_units - 1:
            break
        out += ch
    return out + "…"


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.strip().lstrip("#").upper())


_AFTER_LATIN = ["a:ea", "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst"]
_AFTER_EA = _AFTER_LATIN[1:]


def _ensure_font_el(rpr, tag: str, typeface: str, successors: List[str]) -> None:
    el = rpr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        for s in successors:
            nxt = rpr.find(qn(s))
            if nxt is not None:
                nxt.addprevious(el)
                break
        else:
            rpr.append(el)
    el.set("typeface", typeface)


# 游ゴシック等が無い環境で「ゴシック体（サンセリフ）」に置き換えてもらうためのヒント（Office テーマと同じ値）
GOTHIC_HINT = {"panose": "020B0400000000000000", "pitchFamily": "34", "charset": "-128"}


def _is_gothic(face: str) -> bool:
    return any(k in face for k in ("ゴシック", "Gothic", "gothic", "Meiryo", "メイリオ", "Hiragino", "ヒラギノ", "Sans"))


def set_rpr_fonts(rpr, latin: str, ea: str) -> None:
    """a:rPr / a:defRPr / a:endParaRPr に latin と ea（東アジア）の両方を入れる。"""
    _ensure_font_el(rpr, "a:latin", latin, _AFTER_LATIN)
    _ensure_font_el(rpr, "a:ea", ea, _AFTER_EA)
    for tag, face in (("a:latin", latin), ("a:ea", ea)):
        el = rpr.find(qn(tag))
        if el is not None and _is_gothic(face):
            for k, v in GOTHIC_HINT.items():
                if k == "charset" and tag == "a:latin":
                    continue
                el.set(k, v)
    if rpr.tag == qn("a:rPr") and rpr.get("lang") is None:
        rpr.set("lang", "ja-JP")


def apply_fonts_everywhere(root, latin: str, ea: str) -> None:
    for el in root.iter(qn("a:rPr"), qn("a:defRPr"), qn("a:endParaRPr")):
        set_rpr_fonts(el, latin, ea)


def set_theme_fonts(prs, latin: str, ea: str) -> None:
    """既定テンプレートのテーマフォント（見出し・本文）も指定のフォントにする。"""
    try:
        part = prs.slide_master.part.part_related_by(RT.THEME)
    except KeyError:
        return
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    root = part._element if hasattr(part, "_element") else etree.fromstring(part.blob)
    for tag in ("majorFont", "minorFont"):
        f = root.find(f".//a:fontScheme/a:{tag}", ns)
        if f is None:
            continue
        for sub, face in (("latin", latin), ("ea", ea)):
            el = f.find(f"a:{sub}", ns)
            if el is not None:
                el.set("typeface", face)
        for el in f.findall("a:font", ns):
            if el.get("script") in ("Jpan", "Hans", "Hant", "Hang"):
                if el.get("script") == "Jpan":
                    el.set("typeface", ea)
    if not hasattr(part, "_element"):
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def set_cell_border(cell, bottom_color: str) -> None:
    """セルの枠線: 左右上は無し、下だけ細線（表スタイルの白い罫線に頼らない）。"""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    els = []
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = OxmlElement(tag)
        if tag == "a:lnB":
            ln.set("w", "6350")
            sf = OxmlElement("a:solidFill")
            clr = OxmlElement("a:srgbClr")
            clr.set("val", bottom_color)
            sf.append(clr)
            ln.append(sf)
        else:
            ln.set("w", "0")
            ln.append(OxmlElement("a:noFill"))
        els.append(ln)
    for ln in reversed(els):
        tcPr.insert(0, ln)


def fit_pt(paragraphs: Sequence[str], w_in: float, h_in: float, max_pt: float, min_pt: float,
           spacing: float = 1.35, pad_w: float = 0.2, pad_h: float = 0.12) -> Tuple[float, bool]:
    """箱に収まる最大のフォントサイズを見積もる。戻り値: (pt, 収まったか)"""
    pt = max_pt
    while pt >= min_pt - 1e-9:
        per_line = max(1.0, (w_in - pad_w) * 72.0 / pt)
        lines = sum(max(1, math.ceil(C.text_units(p) / per_line)) for p in paragraphs)
        if lines * pt * spacing / 72.0 + pad_h <= h_in:
            return pt, True
        pt -= 0.5
    return min_pt, False


def lines_needed(text: str, w_in: float, pt: float, pad_w: float = 0.2) -> int:
    """折り返しの行数の見積もり（ビューアによって字幅が違うので 8% 余裕を見る）。"""
    per_line = max(1.0, (w_in - pad_w) * 72.0 / pt)
    return sum(max(1, math.ceil(C.text_units(p) * 1.08 / per_line)) for p in str(text).split("\n"))


class Deck:
    def __init__(self, cfg: dict, template: Optional[Path]) -> None:
        self.cfg = cfg
        self.col = cfg["colors"]
        self.font_latin = cfg["font"]["latin"]
        self.font_ea = cfg["font"]["east_asian"]
        self.warnings: List[str] = []
        if template:
            self.prs = Presentation(str(template))
            sld_ids = self.prs.slides._sldIdLst
            for sld in list(sld_ids):
                self.prs.part.drop_rel(sld.rId)
                sld_ids.remove(sld)
            ratio = self.prs.slide_width / self.prs.slide_height
            if abs(ratio - 16 / 9) > 0.02:
                self.warnings.append(f"テンプレートの縦横比が 16:9 ではありません（{ratio:.2f}）。レイアウトが崩れる可能性があります")
        else:
            self.prs = Presentation()
            self.prs.slide_width = Emu(int(W_IN * 914400))
            self.prs.slide_height = Emu(int(H_IN * 914400))
            set_theme_fonts(self.prs, self.font_latin, self.font_ea)
        self.sx = self.prs.slide_width / (W_IN * 914400)
        self.sy = self.prs.slide_height / (H_IN * 914400)
        self.layout = self._pick_layout()

    def _pick_layout(self):
        layouts = list(self.prs.slide_layouts)
        for lo in layouts:
            if C.nfc(lo.name).strip().lower() in ("blank", "白紙", "空白", "白紙のスライド"):
                return lo
        return min(layouts, key=lambda lo: len(lo.placeholders))

    # 座標（設計上のインチ → EMU）
    def X(self, v: float) -> Emu:
        return Emu(int(round(v * 914400 * self.sx)))

    def Y(self, v: float) -> Emu:
        return Emu(int(round(v * 914400 * self.sy)))

    def new_slide(self):
        slide = self.prs.slides.add_slide(self.layout)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        return slide

    # ---- 図形・文字
    def rect(self, slide, x, y, w, h, fill: Optional[str] = None, line: Optional[str] = None,
             line_w: float = 0.75, shape=MSO_SHAPE.RECTANGLE):
        shp = slide.shapes.add_shape(shape, self.X(x), self.Y(y), self.X(w), self.Y(h))
        shp.shadow.inherit = False
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(line_w)
        else:
            shp.line.fill.background()
        return shp

    def text(self, slide, x, y, w, h, paras, size: float = 12, bold: bool = False, color: Optional[str] = None,
             align: str = "left", anchor: str = "top", spacing: float = 1.15, space_after: float = 0,
             margin: float = 0.05, wrap: bool = True):
        """paras: 文字列（改行で段落）か、段落のリスト。段落は文字列か [(文字列, {bold,color,size}), …]。"""
        tb = slide.shapes.add_textbox(self.X(x), self.Y(y), self.X(w), self.Y(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        for side in ("left", "right"):
            setattr(tf, f"margin_{side}", self.X(margin))
        tf.margin_top = self.Y(0.03)
        tf.margin_bottom = self.Y(0.03)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        if isinstance(paras, str):
            paras = paras.split("\n")
        first = True
        for para in paras:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = spacing
            if space_after:
                p.space_after = Pt(space_after)
            runs = [(para, {})] if isinstance(para, str) else para
            for txt, st in runs:
                r = p.add_run()
                r.text = txt
                f = r.font
                f.size = Pt(st.get("size", size))
                f.bold = st.get("bold", bold)
                f.color.rgb = rgb(st.get("color", color or self.col["text"]))
                f.name = self.font_latin
        return tb

    def frame(self, slide, spec: "Spec", page: int, footer: str) -> None:
        col = self.col
        if spec.group_label:
            self.text(slide, ML, 0.2, 8.0, 0.28, spec.group_label, size=10, bold=True, color=col["accent"])
        title = spec.title
        pt, ok = fit_pt([title], CW, 0.62, self.cfg["font"]["title_pt"], 14, spacing=1.1)
        self.text(slide, ML, 0.44, CW, 0.62, title, size=pt, bold=True, color=col["primary"], anchor="middle")
        self.rect(slide, ML, 1.1, CW, 0.025, fill=col["primary"])
        self.rect(slide, ML, 1.1, 1.2, 0.05, fill=col["accent"])
        self.text(slide, ML, 7.04, 9.0, 0.3, footer, size=8, color=col["subtext"])
        self.text(slide, W_IN - MR - 1.5, 7.04, 1.5, 0.3, str(page), size=9, color=col["subtext"], align="right")

    # ---- 表
    def table(self, slide, x, y, w, header: List[str], rows: List[List[str]], widths: List[float],
              aligns: Optional[List[str]] = None, font_pt: float = 10, row_h: float = 0.3,
              styles: Optional[List[dict]] = None):
        """widths は比率。styles は行ごとの {bold, fill, color, cells: {列番号: color}}。"""
        col = self.col
        n_rows, n_cols = len(rows) + 1, len(header)
        total = sum(widths)
        col_w = [w * r / total for r in widths]
        heights = [max(row_h, 0.28)]
        for row in rows:
            need = max(lines_needed(str(c), cw, font_pt) for c, cw in zip(row, col_w))
            heights.append(max(row_h, need * font_pt * (1.25 if need == 1 else 1.45) / 72 + 0.12))
        gf = slide.shapes.add_table(n_rows, n_cols, self.X(x), self.Y(y), self.X(w), self.Y(sum(heights)))
        tbl = gf.table
        tbl.first_row = True
        tbl.horz_banding = False
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = self.X(cw)
        for r, hgt in enumerate(heights):
            tbl.rows[r].height = self.Y(hgt)
        aligns = aligns or (["left"] + ["right"] * (n_cols - 1))
        for r in range(n_rows):
            st = {} if r == 0 else ((styles or [{}] * len(rows))[r - 1] or {})
            for c in range(n_cols):
                cell = tbl.cell(r, c)
                cell.margin_left = self.X(0.06)
                cell.margin_right = self.X(0.06)
                cell.margin_top = self.Y(0.02)
                cell.margin_bottom = self.Y(0.02)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                set_cell_border(cell, col["border"] if r > 0 else col["table_header"])
                cell.fill.solid()
                if r == 0:
                    cell.fill.fore_color.rgb = rgb(col["table_header"])
                else:
                    cell.fill.fore_color.rgb = rgb(st.get("fill") or (col["table_band"] if r % 2 == 0 else "FFFFFF"))
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER if r == 0 else \
                    {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[aligns[c]]
                run = p.add_run()
                run.text = str(header[c] if r == 0 else rows[r - 1][c])
                f = run.font
                f.size = Pt(font_pt - (0.5 if r == 0 else 0))
                f.name = self.font_latin
                f.bold = True if r == 0 else bool(st.get("bold"))
                color = "FFFFFF" if r == 0 else (st.get("cells", {}).get(c) or st.get("color") or col["text"])
                f.color.rgb = rgb(color)
        return sum(heights)

    # ---- グラフ
    def chart(self, slide, kind: str, x, y, w, h, categories: List[str], series: List[Tuple[str, list]],
              number_format: str = "#,##0", legend: bool = True, labels: bool = False,
              label_format: Optional[str] = None, colors: Optional[List[str]] = None, font_pt: float = 9,
              gap: int = 60, overlap: Optional[int] = None):
        """kind: column / bar / line / pie。単一種別のグラフだけを作る（複合グラフは作らない）。"""
        ctype = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                 "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}[kind]
        colors = colors or self.col["series"]
        cd = CategoryChartData(number_format=number_format)
        cd.categories = categories
        digits = 0 if (number_format.startswith('"') or number_format == "#,##0") else \
            (2 if number_format in ("0.00", PCT_FMT) else 4)
        for name, vals in series:
            # 表示の桁に丸めて渡す（書式を無視するビューアでも桁が暴れない。元の値は analysis.json）
            cd.add_series(name, [None if v is None else round(v, digits) for v in vals])
        gf = slide.shapes.add_chart(ctype, self.X(x), self.Y(y), self.X(w), self.Y(h), cd)
        ch = gf.chart
        ch.font.size = Pt(font_pt)
        ch.font.name = self.font_latin
        ch.font.color.rgb = rgb(self.col["text"])
        ch.has_title = False
        show_legend = legend and (len(series) > 1 or kind == "pie")
        ch.has_legend = show_legend
        if show_legend:
            ch.legend.position = XL_LEGEND_POSITION.RIGHT if kind == "pie" else XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(font_pt)
        plot = ch.plots[0]
        if kind != "pie":
            va = ch.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = rgb(self.col["border"])
            va.major_gridlines.format.line.width = Pt(0.5)
            va.tick_labels.number_format = number_format
            va.tick_labels.number_format_is_linked = False
            va.tick_labels.font.size = Pt(font_pt - 1)
            va.major_tick_mark = XL_TICK_MARK.NONE
            va.format.line.fill.background()
            vals = [v for _, vs in series for v in vs if v is not None]
            vmax = max(vals) if vals else 0
            vmin = min(vals) if vals else 0
            if vmin >= 0:
                va.minimum_scale = 0
                room = 1.3 if (labels and kind == "bar") else (1.18 if labels else 1.12)  # データラベルの置き場
                top = nice_max(vmax * room) if vmax > 0 else 1.0
                va.maximum_scale = top
                unit = top / 4
                va.major_unit = None if (number_format == "#,##0" and unit != int(unit)) else unit
            ca = ch.category_axis
            ca.tick_labels.font.size = Pt(font_pt - 1)
            ca.major_tick_mark = XL_TICK_MARK.NONE
            ca.format.line.color.rgb = rgb(self.col["border"])
            if len(categories) > 16:
                set_tick_skip(ch, math.ceil(len(categories) / 15))
            if kind == "bar":
                longest = max((C.text_units(str(c)) for c in categories), default=4)
                lab_w = min(0.45, max(0.12, (longest * (font_pt - 1) / 72 + 0.15) / w))
                right = 0.1 if labels else 0.04
                top_pad = 0.1 if show_legend else 0.03
                set_plot_layout(ch, lab_w, top_pad, max(0.3, 1 - lab_w - right), 0.97 - top_pad - 0.08)
        if kind in ("column", "bar"):
            plot.gap_width = gap
            if overlap is not None:
                plot.overlap = overlap
        if kind == "pie":
            for i, pt in enumerate(plot.series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
        else:
            for i, s in enumerate(plot.series):
                c = rgb(colors[i % len(colors)])
                if kind == "line":
                    s.format.line.color.rgb = c
                    s.format.line.width = Pt(2)
                    s.smooth = False
                    s.marker.style = XL_MARKER_STYLE.CIRCLE
                    s.marker.size = 5
                    s.marker.format.fill.solid()
                    s.marker.format.fill.fore_color.rgb = c
                    s.marker.format.line.color.rgb = c
                else:
                    s.format.fill.solid()
                    s.format.fill.fore_color.rgb = c
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(font_pt - 1)
            if kind == "pie":
                dl.show_percentage = True
                dl.show_value = False
                dl.show_category_name = False
                dl.number_format = "0%"
                dl.number_format_is_linked = False
            else:
                dl.number_format = label_format or number_format
                dl.number_format_is_linked = False
                dl.position = XL_LABEL_POSITION.ABOVE if kind == "line" else XL_LABEL_POSITION.OUTSIDE_END
        return ch

    def chart_title(self, slide, x, y, w, text: str) -> None:
        self.text(slide, x, y, w, 0.3, text, size=10.5, bold=True, color=self.col["text"])


# ---------------------------------------------------------------- スライド計画

class Spec:
    def __init__(self, sid: str, title: str, render: Optional[Callable], section: str,
                 scopes: Sequence[str] = (), keys: str = "", slot: Optional[str] = None, kind: str = "content",
                 extra_notes: Sequence[str] = ()) -> None:
        self.id = sid
        self.title = title
        self.render = render
        self.section = section
        self.scopes = list(scopes)
        self.keys = keys
        self.slot = slot
        self.kind = kind
        self.extra_notes = list(extra_notes)
        self.group = ""
        self.group_label = ""
        self.page = 0


class Builder:
    def __init__(self, analysis: dict, insights: dict, cfg: dict, deck: Deck, client_dir: Path, vault: Path,
                 insights_path: Optional[Path]) -> None:
        self.a = analysis
        self.v: Dict[str, dict] = analysis["values"]
        self.cfg = cfg
        self.deck = deck
        self.col = cfg["colors"]
        self.client_dir = client_dir
        self.vault = vault
        self.defs = analysis["definitions"]
        self.slots = {s["id"]: s for s in analysis["insight_slots"]}
        self.insights = insights
        self.insights_path = insights_path
        self.resolved: Dict[str, Tuple[str, List[str]]] = {}
        self.unknown_count = 0
        self.specs: List[Spec] = []
        cur = analysis.get("currency", "JPY")
        sym = {"JPY": "¥", "USD": "$", "EUR": "€"}.get(cur, "")
        self.fmt_cur = f'"{sym}"#,##0' if cur == "JPY" else f'"{sym}"#,##0.00'
        self.camp_by_id = {c["id"]: c for c in analysis["campaigns"]}
        self.adset_by_id = {s["id"]: s for s in analysis["adsets"]}
        self.ad_by_id = {a["id"]: a for a in analysis["ads"]}

    # ---- 所見
    def resolve_insights(self) -> List[str]:
        """全スロットのプレースホルダを解決する。戻り値: エラーの一覧（空なら OK）。"""
        errors: List[str] = []
        slides = (self.insights or {}).get("slides", {}) if isinstance(self.insights, dict) else {}
        if self.insights and not isinstance(slides, dict):
            return ["insights.json の slides がオブジェクト（{ }）ではありません"]
        for sid, slot in slides.items():
            text = C.insight_text(slot)
            if not text:
                continue
            resolved, unknown = C.resolve_placeholders(text, self.v)
            for k in unknown:
                errors.append(f"slides[\"{sid}\"]: 存在しないキー {{{{{k}}}}}")
            if sid not in self.slots:
                self.deck.warnings.append(f"insights.json の slides[\"{sid}\"] に対応するスライドがありません（表示されません）")
            self.resolved[sid] = (resolved, C.insight_sources(slot))
        return errors

    def insight(self, slide, x, y, w, h, slot_id: str, title: str = "所見") -> None:
        d, col = self.deck, self.col
        text, sources = self.resolved.get(slot_id, ("", []))
        if text:
            d.rect(slide, x, y, w, h, fill=col["accent_light"])
            d.rect(slide, x, y, 0.06, h, fill=col["accent"])
            d.text(slide, x + 0.15, y + 0.1, w - 0.25, 0.32, title, size=11, bold=True, color=col["accent"])
            paras = [p for p in text.split("\n") if p.strip()]
            pt, ok = fit_pt(paras, w - 0.3, h - 0.55, self.cfg["font"]["body_pt"], 8)
            if not ok:
                d.warnings.append(f"所見「{slot_id}」が長く、枠からはみ出す可能性があります（{len(text)}文字）")
            d.text(slide, x + 0.15, y + 0.45, w - 0.3, h - 0.55, paras, size=pt, color=col["text"], spacing=1.3,
                   space_after=3)
        else:
            self.unknown_count += 1
            d.rect(slide, x, y, w, h, fill=col["warn_fill"], line=col["warn_border"], line_w=1.5)
            d.text(slide, x + 0.15, y + 0.12, w - 0.3, 0.45, C.UNKNOWN_INSIGHT, size=15, bold=True,
                   color=col["warn_text"])
            lines = ["所見が未記入です（数字は表・グラフを参照）。"]
            cand = self.slots.get(slot_id, {}).get("sources") or []
            if cand:
                lines.append("参照できる記録:")
                lines += ["・" + c for c in cand]
            pt, _ = fit_pt(lines, w - 0.3, h - 0.7, 10, 7)
            d.text(slide, x + 0.15, y + 0.62, w - 0.3, h - 0.72, lines, size=pt, color=col["warn_text"])

    # ---- 値の取り出し
    def disp(self, m: Optional[dict], k: str, mark_ref: bool = True) -> str:
        if not m or k not in m:
            return C.NA_DISPLAY
        e = m[k]
        s = e.get("display", C.NA_DISPLAY)
        return s + ("※" if mark_ref and e.get("reference") else "")

    def val(self, m: Optional[dict], k: str):
        if not m or k not in m:
            return None
        return m[k].get("value")

    def label(self, k: str) -> str:
        return self.defs.get(k, {}).get("label", k)

    def unit_fmt(self, k: str) -> str:
        u = self.defs.get(k, {}).get("unit")
        return {"currency": self.fmt_cur, "ratio": "0.00%", "x": "0.00", "freq": "0.00"}.get(u, "#,##0")

    def trend_color(self, trend: str) -> str:
        return {"改善": self.col["good"], "悪化": self.col["bad"]}.get(trend, self.col["subtext"])

    def short_dates(self) -> List[str]:
        """日別グラフのカテゴリ（日にちだけ。月はグラフのタイトルに書く）。"""
        return [str(int(d[8:10])) for d in self.a["daily_series"]["dates"]]

    def mlabel(self) -> str:
        return f"（{int(self.a['month'][5:7])}月）"

    # ---- 部品
    def kpi_cards(self, slide, x, y, w, h, keys: List[str], m: dict, mom: Optional[dict], cols: int = 4,
                  gap: float = 0.16, targets: Optional[Dict[str, dict]] = None, value_pt: float = 20) -> None:
        d, col = self.deck, self.col
        keys = [k for k in keys if k in m] or keys
        rows = math.ceil(len(keys) / cols)
        cw = (w - gap * (cols - 1)) / cols
        chh = (h - gap * (rows - 1)) / rows
        for i, k in enumerate(keys):
            cx = x + (i % cols) * (cw + gap)
            cy = y + (i // cols) * (chh + gap)
            d.rect(slide, cx, cy, cw, chh, fill="FFFFFF", line=col["border"], line_w=0.75)
            d.rect(slide, cx, cy, cw, 0.05, fill=col["accent"])
            e = m.get(k)
            lab = self.label(k) + ("（参考値）" if e and e.get("reference") else "")
            d.text(slide, cx + 0.08, cy + 0.1, cw - 0.16, 0.3, lab, size=9.5, color=col["subtext"])
            vs = e.get("display", C.NA_DISPLAY) if e else C.NA_DISPLAY
            pt, _ = fit_pt([vs], cw - 0.1, 0.55, value_pt, 11, spacing=1.0, pad_h=0)
            d.text(slide, cx + 0.08, cy + 0.36, cw - 0.16, chh * 0.42, vs, size=pt, bold=True,
                   color=col["primary"], anchor="middle")
            subs = []
            if mom and k in mom:
                tr = mom[k]["trend"]["display"]
                subs.append([(f"前月比 {mom[k]['pct']['display']}", {"color": self.trend_color(tr)})])
            if targets and k in targets:
                t = targets[k]
                c2 = col["good"] if t["status"] in ("達成", "予算内") else col["bad"]
                subs.append([(f"目標 {t['goal']['display']}・{t['status']}", {"color": c2})])
            if subs:
                sh = 0.2 * len(subs) + 0.08
                d.text(slide, cx + 0.08, cy + chh - sh - 0.02, cw - 0.16, sh, subs, size=8.5, spacing=1.0)

    def paginate(self, rows: List[List[str]], widths_in: List[float], font_pt: float, max_rows: int,
                 avail_h: float, row_h: float = 0.3) -> List[List[List[str]]]:
        pages, cur, used = [], [], max(row_h, 0.28)
        for row in rows:
            need = max(lines_needed(str(c), cw, font_pt) for c, cw in zip(row, widths_in))
            hh = max(row_h, need * font_pt * (1.25 if need == 1 else 1.45) / 72 + 0.12) * 1.1  # ビューア差を見込む
            if cur and (len(cur) >= max_rows or used + hh > avail_h):
                pages.append(cur)
                cur, used = [], max(row_h, 0.28)
            cur.append(row)
            used += hh
        if cur:
            pages.append(cur)
        return pages or [[]]

    def paged_table_specs(self, sid: str, title: str, section: str, header: List[str], rows: List[List[str]],
                          widths: List[float], aligns: Optional[List[str]], max_rows: int, scopes: Sequence[str],
                          keys: str, font_pt: Optional[float] = None, styles: Optional[List[dict]] = None,
                          top_note: str = "") -> List[Spec]:
        font_pt = font_pt or self.cfg["font"]["table_pt"]
        total = sum(widths)
        widths_in = [CW * r / total for r in widths]
        top = CT + (0.35 if top_note else 0)
        pages = self.paginate(rows, widths_in, font_pt, max_rows, CB - top)
        specs = []
        offset = 0
        for i, page_rows in enumerate(pages, start=1):
            st = styles[offset:offset + len(page_rows)] if styles else None
            offset += len(page_rows)
            t = title + (f"（{i}/{len(pages)}）" if len(pages) > 1 else "")

            def render(b, slide, spec, page_rows=page_rows, st=st):
                if top_note:
                    b.deck.text(slide, ML, CT - 0.02, CW, 0.3, top_note, size=9, color=b.col["subtext"])
                b.deck.table(slide, ML, top, CW, header, page_rows, widths, aligns, font_pt=font_pt, row_h=0.3,
                             styles=st)
            specs.append(Spec(f"{sid}.p{i}" if len(pages) > 1 else sid, t, render, section, scopes, keys))
        return specs

    # ---------------------------------------------------------------- 各セクション
    def sec_cover(self) -> List[Spec]:
        def render(b, slide, spec):
            d, col, cfg = b.deck, b.col, b.cfg["cover"]
            d.rect(slide, 0, 0, W_IN, H_IN, fill=col["primary"])
            d.rect(slide, 0.9, 2.05, 0.12, 2.3, fill=col["accent"])
            d.text(slide, 1.25, 1.9, 11, 0.9, cfg.get("title", "広告運用レポート"), size=36, bold=True, color="FFFFFF")
            d.text(slide, 1.25, 2.8, 11, 0.5, cfg.get("subtitle", ""), size=16, color="C9D6EA")
            d.text(slide, 1.25, 3.55, 11, 0.7, f"{b.a['client']} 様", size=26, bold=True, color="FFFFFF")
            d.text(slide, 1.25, 4.6, 11, 0.4, f"対象期間：{b.a['period']['label']}", size=14, color="FFFFFF")
            d.text(slide, 1.25, 5.0, 11, 0.4, f"作成日：{dt.date.today().year}年{dt.date.today().month}月"
                   f"{dt.date.today().day}日", size=14, color="FFFFFF")
            if cfg.get("author"):
                d.text(slide, 1.25, 5.4, 11, 0.4, cfg["author"], size=14, color="FFFFFF")
            d.text(slide, 1.25, 6.55, 11.5, 0.5,
                   "本資料の数値は広告マネージャのエクスポートから自動集計しています（各ページのノートに出どころを記載）。",
                   size=10, color="C9D6EA")
        return [Spec("cover", "表紙", render, "cover", kind="cover")]

    def sec_toc(self) -> List[Spec]:
        def render(b, slide, spec):
            d, col = b.deck, b.col
            groups: List[Tuple[str, str, int, List[Tuple[str, int]]]] = []
            for s in b.specs:
                if s.kind in ("cover", "toc") or not s.group:
                    continue
                if not groups or groups[-1][0] != s.group:
                    groups.append((s.group, s.group_label, s.page, []))
                if s.kind == "content":
                    titles = [t for t, _ in groups[-1][3]]
                    sec_title = SECTION_TITLES.get(s.section, s.section)
                    if sec_title not in titles:
                        groups[-1][3].append((sec_title, s.page))
            y = CT + 0.1
            n = len(groups)
            row_h = min(0.7, (CH - 0.1) / max(n, 1))
            for i, (_, glabel, page, secs) in enumerate(groups, start=1):
                d.rect(slide, ML, y + 0.08, 0.42, row_h - 0.2, fill=col["primary"])
                d.text(slide, ML, y + 0.08, 0.42, row_h - 0.2, str(i) if glabel != "Appendix" else "A", size=14,
                       bold=True, color="FFFFFF", align="center", anchor="middle")
                d.text(slide, ML + 0.6, y + 0.02, 5.0, 0.4, glabel, size=15, bold=True, color=col["primary"])
                d.text(slide, ML + 0.6, y + 0.36, 10.6, 0.32,
                       "　/　".join(f"{t}（p.{p}）" for t, p in secs), size=9, color=col["subtext"])
                d.text(slide, W_IN - MR - 1.2, y + 0.05, 1.2, 0.4, f"p.{page}", size=13, bold=True,
                       color=col["accent"], align="right")
                y += row_h
        return [Spec("toc", "目次", render, "toc", kind="toc")]

    def sec_summary(self) -> List[Spec]:
        def render(b, slide, spec):
            a = b.a
            targets = {t["metric"]: t for t in a["targets"]["items"]}
            b.kpi_cards(slide, ML, CT, LEFT_W, 3.1, b.cfg["kpi_cards"]["summary"], a["totals"], a["mom"],
                        targets=targets)
            y = CT + 3.3
            b.deck.chart_title(slide, ML, y, LEFT_W, "キャンペーン別の概況")
            rows, styles = [], []
            for c in a["campaigns"][:6]:
                m = c["metrics"]
                rows.append([c["name"], c["result_label"], b.disp(m, "spend"), b.disp(m, "cv"), b.disp(m, "cpa"),
                             b.disp(m, "roas"), b.disp(m, "results")])
            b.deck.table(slide, ML, y + 0.32, LEFT_W, ["キャンペーン", "結果の種類", "消化金額", "CV", "CPA", "ROAS", "結果"],
                         rows, [2.1, 2.0, 1.3, 0.7, 1.0, 0.9, 1.3], ["left", "left"] + ["right"] * 5, font_pt=9.5,
                         row_h=0.3)
            if len(a["campaigns"]) > 6:
                b.deck.text(slide, ML, CB - 0.25, LEFT_W, 0.25, f"※ 消化金額の上位6件を表示（全{len(a['campaigns'])}件）",
                            size=8, color=b.col["subtext"])
            b.insight(slide, INS_X, CT, INS_W, CH, "summary")
        return [Spec("summary", "エグゼクティブサマリー", render, "summary", ["total", "prev_total"],
                     "values の total.* / mom.* / target.* / campaign.*", slot="summary",
                     extra_notes=["※ カードの「前月比」は mom.<指標>.pct、色は改善=緑・悪化=赤（指標の良し悪しの向きで判定）"])]

    def sec_targets(self) -> List[Spec]:
        def render(b, slide, spec):
            a, d, col = b.a, b.deck, b.col
            items = a["targets"]["items"]
            if not items:
                d.rect(slide, ML, CT + 0.3, LEFT_W, 2.0, fill=col["bg_light"], line=col["border"])
                d.text(slide, ML + 0.3, CT + 0.5, LEFT_W - 0.6, 0.6, "目標未設定", size=24, bold=True,
                       color=col["subtext"])
                d.text(slide, ML + 0.3, CT + 1.2, LEFT_W - 0.6, 0.9, "今月は広告のKPI目標が設定されていません。",
                       size=13, color=col["subtext"])
            else:
                rows, styles = [], []
                for t in items:
                    rows.append([t["label"], t["goal"]["display"], t["actual"]["display"],
                                 t["achievement"]["display"], t["status"]])
                    good = t["status"] in ("達成", "予算内")
                    styles.append({"cells": {4: col["good"] if good else col["bad"]}, "bold": False})
                h = d.table(slide, ML, CT, LEFT_W, ["指標", "目標", "実績", "達成率", "判定"], rows,
                            [2.2, 1.5, 1.5, 1.3, 1.1], ["left", "right", "right", "right", "center"], font_pt=11,
                            row_h=0.4, styles=styles)
                cy = CT + h + 0.3
                if CB - cy > 1.6:
                    d.chart_title(slide, ML, cy, LEFT_W, "達成率（100%＝目標どおり。CPA等は低いほど達成率が高い）")
                    vals = [pct100(t["achievement"]["value"]) for t in items][::-1]
                    ch = d.chart(slide, "bar", ML, cy + 0.3, LEFT_W, CB - cy - 0.3, [t["label"] for t in items][::-1],
                                 [("達成率（%）", vals)], number_format=PCT_FMT, legend=False, labels=True)
                    ch.value_axis.maximum_scale = max(120.0, math.ceil(max([v or 0 for v in vals]) / 20.0) * 20 + 20)
                    ch.value_axis.major_unit = 20
            b.insight(slide, INS_X, CT, INS_W, CH, "targets")
        notes = [] if self.a["targets"]["items"] else [
            "目標を表示するには、クライアントの _profile.md の frontmatter に "
            "`ad_kpi_targets: [cpa=3000, roas=3.0]` のように書いて ad_analyze.py を再実行する。"]
        return [Spec("targets", "目標対比", render, "targets", ["total"] + [f"target.{t['metric']}" for t in
                                                                     self.a["targets"]["items"]],
                     "values の target.<指標>.goal / actual / achievement / status（目標は _profile.md の ad_kpi_targets）",
                     slot="targets", extra_notes=notes)]

    def sec_mom(self) -> List[Spec]:
        a = self.a
        if not a["has_prev"]:
            def render_none(b, slide, spec):
                d, col = b.deck, b.col
                d.rect(slide, ML, CT + 0.3, LEFT_W, 2.0, fill=col["bg_light"], line=col["border"])
                d.text(slide, ML + 0.3, CT + 0.5, LEFT_W - 0.6, 0.6, "前月データなし", size=24, bold=True,
                       color=col["subtext"])
                d.text(slide, ML + 0.3, CT + 1.2, LEFT_W - 0.6, 0.9, "前月の広告データが無いため、前月比較は表示していません。",
                       size=13, color=col["subtext"])
                b.insight(slide, INS_X, CT, INS_W, CH, "mom")
            return [Spec("mom", "前月比較", render_none, "mom", ["total"], "（前月データなし）", slot="mom",
                         extra_notes=["前月比を出すには、前月のエクスポートを 01_clients/<client>/ads/<前月>/ に置いて ad_analyze.py を再実行する。"])]

        keys = [k for k in ["spend", "impressions", "reach", "frequency", "link_clicks", "ctr", "cpc", "cpm", "cv",
                            "cvr", "cpa", "purchase_value", "roas"] if k in a["totals"]]

        def render(b, slide, spec):
            rows, styles = [], []
            for k in keys:
                mo = (a["mom"] or {}).get(k)
                tr = mo["trend"]["display"] if mo else C.NA_DISPLAY
                rows.append([b.label(k) + ("※" if a["totals"][k].get("reference") else ""),
                             b.disp(a["totals"], k, False), b.disp(a["prev_totals"], k, False),
                             mo["diff"]["display"] if mo else C.NA_DISPLAY,
                             mo["pct"]["display"] if mo else C.NA_DISPLAY, tr])
                styles.append({"cells": {5: b.trend_color(tr), 4: b.trend_color(tr)}})
            b.deck.table(slide, ML, CT, LEFT_W, ["指標", b.v["month.label"]["display"], b.v["prev_month.label"]["display"],
                                                 "差分", "増減率", "評価"], rows, [2.0, 1.5, 1.5, 1.3, 1.0, 0.8],
                         font_pt=10, row_h=0.36, styles=styles)
            if any(a["totals"][k].get("reference") for k in keys):
                b.deck.text(slide, ML, CB - 0.28, LEFT_W, 0.28, "※ リーチ・フリークエンシーは日別合算の参考値（重複を含む）",
                            size=8, color=b.col["subtext"])
            b.insight(slide, INS_X, CT, INS_W, CH, "mom")

        def render2(b, slide, spec):
            d = b.deck
            camps = [c for c in a["campaigns"]][:8]
            cats = [short_label(c["name"], 12) for c in camps]
            hw = (CW - 0.3) / 2
            d.chart_title(slide, ML, CT, hw, "消化金額（前月・今月）")
            d.chart(slide, "column", ML, CT + 0.3, hw, 2.75, cats,
                    [("前月", [b.val(c["prev_metrics"], "spend") for c in camps]),
                     ("今月", [b.val(c["metrics"], "spend") for c in camps])], number_format=b.fmt_cur,
                    colors=[b.col["series"][2], b.col["series"][0]])
            d.chart_title(slide, ML + hw + 0.3, CT, hw, "CPA（前月・今月）")
            d.chart(slide, "column", ML + hw + 0.3, CT + 0.3, hw, 2.75, cats,
                    [("前月", [b.val(c["prev_metrics"], "cpa") for c in camps]),
                     ("今月", [b.val(c["metrics"], "cpa") for c in camps])], number_format=b.fmt_cur,
                    colors=[b.col["series"][2], b.col["series"][0]])
            rows, styles = [], []
            for c in a["campaigns"][:8]:
                mo = c["mom"] or {}
                rows.append([c["name"], b.disp(c["metrics"], "spend"), b.disp(c["prev_metrics"], "spend"),
                             (mo.get("spend") or {}).get("pct", {}).get("display", C.NA_DISPLAY),
                             b.disp(c["metrics"], "cpa"), b.disp(c["prev_metrics"], "cpa"),
                             (mo.get("cpa") or {}).get("pct", {}).get("display", C.NA_DISPLAY)])
                tr = (mo.get("cpa") or {}).get("trend", {}).get("display", "")
                styles.append({"cells": {6: b.trend_color(tr)}})
            d.table(slide, ML, CT + 3.25, CW, ["キャンペーン", "消化金額（今月）", "（前月）", "増減率", "CPA（今月）", "（前月）", "増減率"],
                    rows, [3.0, 1.5, 1.4, 1.0, 1.4, 1.3, 1.0], font_pt=9.5, row_h=0.3, styles=styles)
        return [Spec("mom", "前月比較（全体）", render, "mom", ["total", "prev_total"],
                     "values の total.* / prev.* / mom.<指標>.diff・pct・trend", slot="mom"),
                Spec("mom.campaigns", "前月比較（キャンペーン別）", render2, "mom",
                     [f"campaign.{c['id']}" for c in a["campaigns"]],
                     "values の campaign.<ID>.* / campaign.<ID>.prev.* / campaign.<ID>.mom.*（前月は同名キャンペーンで照合）")]

    def sec_measurement_notes(self) -> List[Spec]:
        a = self.a
        notes_cfg = self.cfg.get("measurement_notes", [])
        prev_start = dt.date.fromisoformat(a["prev_month"] + "-01")
        end = dt.date.fromisoformat(a["period"]["end"])

        def render(b, slide, spec):
            d, col = b.deck, b.col
            y = CT
            d.chart_title(slide, ML, y, CW, "アトリビューション設定（エクスポートの値）")
            attr = a["attribution"]
            if attr["available"]:
                rows = []
                for c in a["campaigns"]:
                    vals = attr["by_campaign"].get(c["id"], [])
                    rows.append([c["name"], " / ".join(vals) if vals else "（値なし）"])
                h = d.table(slide, ML, y + 0.32, CW * 0.62, ["キャンペーン", "アトリビューション設定"], rows, [1.2, 1.6],
                            aligns=["left", "left"], font_pt=9.5, row_h=0.28)
            else:
                d.text(slide, ML, y + 0.32, CW, 0.3, "エクスポートに「アトリビューション設定」の列がありません（要確認）。",
                       size=10, color=col["warn_text"])
                h = 0.3
            y = y + 0.32 + h + 0.2
            d.chart_title(slide, ML, y, CW, "Meta の計測仕様の変更（前月比に段差が出うるもの）")
            y += 0.34
            for n in notes_cfg:
                try:
                    nd = dt.date.fromisoformat(n.get("date", ""))
                    crosses = prev_start <= nd <= end
                except ValueError:
                    crosses = False
                flag = ("　▶ 今回の比較期間（" + b.v["prev_month.label"]["display"] + "→" + b.v["month.label"]["display"] +
                        "）はこの変更日をまたぎます。前月比の差に注意してください。") if crosses else ""
                paras = [[(n.get("text", ""), {})]]
                if flag:
                    paras.append([(flag.strip(), {"bold": True, "color": col["bad"]})])
                paras.append([("出典：" + n.get("source", "") + "　" + n.get("url", ""),
                               {"size": 8, "color": col["subtext"]})])
                need = sum(lines_needed(p[0][0], CW - 0.4, 9.5) for p in paras) * 9.5 * 1.3 / 72 + 0.15
                d.rect(slide, ML, y, 0.05, need, fill=col["accent"])
                d.text(slide, ML + 0.12, y, CW - 0.12, need, paras, size=9.5, spacing=1.2)
                y += need + 0.12
            d.chart_title(slide, ML, y, CW, "本レポートの集計ルール")
            rules = [
                f"・CV = {a['definitions']['cv']['formula']}。CPA = 消化金額 ÷ CV、CVR = CV ÷ リンククリック（分母はリンククリック）。",
                "・CTR・CPC・CPM・CPA・ROAS・CVR は、エクスポートの比率の列を使わず、元の数値（消化金額・表示回数・クリック・CV）から計算し直しています。",
                "・「結果」はキャンペーンの目的ごとに意味が違う（結果インジケーター）ため、目的の違うキャンペーンの結果は合算していません。",
                "・リーチは足し算できないため、期間合計のエクスポートがある範囲だけ実数、それ以外は日別合算の参考値（重複を含む）です（表の※）。",
            ]
            pt, _ = fit_pt(rules, CW, CB - y - 0.3, 9.5, 8)
            d.text(slide, ML, y + 0.3, CW, CB - y - 0.3, rules, size=pt, spacing=1.2, color=col["text"])
        return [Spec("measurement_notes", "計測上の注意", render, "measurement_notes",
                     [f"campaign.{c['id']}" for c in a["campaigns"]],
                     "analysis.json の attribution（アトリビューション設定の列）と deck_config.json の measurement_notes（固定文言）",
                     extra_notes=["計測仕様の変更の文言は deck_config.json の measurement_notes（出典URL付き）。"])]

    def sec_daily(self) -> List[Spec]:
        a = self.a
        ser = a["daily_series"]
        cats = self.short_dates()
        specs: List[Spec] = []

        def r1(b, slide, spec):
            d = b.deck
            d.chart_title(slide, ML, CT, LEFT_W, "日別 消化金額" + b.mlabel())
            d.chart(slide, "column", ML, CT + 0.3, LEFT_W, 2.55, cats, [("消化金額", ser["spend"])],
                    number_format=b.fmt_cur, legend=False, gap=40)
            d.chart_title(slide, ML, CT + 2.95, LEFT_W, f"日別 {b.label('cv')}" + b.mlabel())
            d.chart(slide, "column", ML, CT + 3.25, LEFT_W, 2.3, cats, [(b.label("cv"), ser["cv"])],
                    number_format="#,##0", legend=False, gap=40, colors=[b.col["series"][1]])
            b.insight(slide, INS_X, CT, INS_W, CH, "daily")

        def r2(b, slide, spec):
            d = b.deck
            d.chart_title(slide, ML, CT, CW, "日別 CPA（CVが0件の日は空白）" + b.mlabel())
            d.chart(slide, "line", ML, CT + 0.3, CW, 2.45, cats, [("CPA", ser["cpa"])], number_format=b.fmt_cur,
                    legend=False, colors=[b.col["series"][3]])
            d.chart_title(slide, ML, CT + 2.85, CW, "日別 ROAS" + b.mlabel())
            d.chart(slide, "line", ML, CT + 3.15, CW, 2.4, cats, [("ROAS", ser["roas"])], number_format="0.00",
                    legend=False, colors=[b.col["series"][4]])

        def r3(b, slide, spec):
            d = b.deck
            d.chart_title(slide, ML, CT, CW, "日別 CTR（リンク）と CVR" + b.mlabel())
            d.chart(slide, "line", ML, CT + 0.3, CW, 2.45, cats,
                    [("CTR（リンク）%", [pct100(v) for v in ser["ctr"]]), ("CVR %", [pct100(v) for v in ser["cvr"]])],
                    number_format=PCT_FMT, colors=[b.col["series"][1], b.col["series"][3]])
            cpc = [m["metrics"].get("cpc", {}).get("value") for m in a["daily"]]
            cpm = [m["metrics"].get("cpm", {}).get("value") for m in a["daily"]]
            hw = (CW - 0.3) / 2
            d.chart_title(slide, ML, CT + 2.85, hw, "日別 CPC（リンク）" + b.mlabel())
            d.chart(slide, "line", ML, CT + 3.15, hw, 2.4, cats, [("CPC", cpc)], number_format=b.fmt_cur,
                    legend=False, colors=[b.col["series"][0]])
            d.chart_title(slide, ML + hw + 0.3, CT + 2.85, hw, "日別 CPM" + b.mlabel())
            d.chart(slide, "line", ML + hw + 0.3, CT + 3.15, hw, 2.4, cats, [("CPM", cpm)], number_format=b.fmt_cur,
                    legend=False, colors=[b.col["series"][5]])

        specs.append(Spec("daily", "日別推移（消化金額・CV）", r1, "daily", ["total"],
                          "analysis.json の daily_series / values の day.<日付>.*", slot="daily"))
        specs.append(Spec("daily.cpa", "日別推移（CPA・ROAS）", r2, "daily", ["total"],
                          "analysis.json の daily_series.cpa / roas（日ごとに 消化金額÷CV、購入金額÷消化金額）"))
        specs.append(Spec("daily.rates", "日別推移（CTR・CVR・CPC・CPM）", r3, "daily", ["total"],
                          "analysis.json の daily_series / daily[].metrics（日ごとに元の数値から再計算）"))
        rows = []
        for dd in a["daily"]:
            m = dd["metrics"]
            rows.append([dd["label"], self.disp(m, "spend"), self.disp(m, "impressions"), self.disp(m, "link_clicks"),
                         self.disp(m, "ctr"), self.disp(m, "cpc"), self.disp(m, "cpm"), self.disp(m, "cv"),
                         self.disp(m, "cvr"), self.disp(m, "cpa"), self.disp(m, "roas")])
        specs += self.paged_table_specs("daily.table", "日別明細", "daily",
                                        ["日付", "消化金額", "IMP", "リンククリック", "CTR", "CPC", "CPM", "CV", "CVR", "CPA",
                                         "ROAS"], rows, [1.1, 1.3, 1.2, 1.1, 0.9, 0.9, 0.9, 0.8, 0.9, 1.0, 0.9], None,
                                        self.cfg["rows_per_page"].get("daily_table", 16), ["total"],
                                        "values の day.<日付>.*")
        return specs

    def sec_weekly(self) -> List[Spec]:
        a = self.a

        def r1(b, slide, spec):
            d = b.deck
            wk = a["weekly"]
            cats = [w["id"] for w in wk]
            hw = (LEFT_W - 0.3) / 2
            d.chart_title(slide, ML, CT, hw, "週別 消化金額")
            d.chart(slide, "column", ML, CT + 0.3, hw, 2.4, cats, [("消化金額", [b.val(w["metrics"], "spend") for w in wk])],
                    number_format=b.fmt_cur, legend=False)
            d.chart_title(slide, ML + hw + 0.3, CT, hw, "週別 CPA")
            d.chart(slide, "column", ML + hw + 0.3, CT + 0.3, hw, 2.4, cats,
                    [("CPA", [b.val(w["metrics"], "cpa") for w in wk])], number_format=b.fmt_cur, legend=False,
                    colors=[b.col["series"][3]])
            rows = [[w["id"], w["label"].split(" ", 1)[1], f"{w['days']}日", b.disp(w["metrics"], "spend"),
                     b.disp(w["metrics"], "cv"), b.disp(w["metrics"], "cpa"), b.disp(w["metrics"], "ctr"),
                     b.disp(w["metrics"], "roas")] for w in wk]
            d.table(slide, ML, CT + 2.95, LEFT_W, ["週", "期間", "日数", "消化金額", "CV", "CPA", "CTR", "ROAS"], rows,
                    [0.6, 1.5, 0.7, 1.4, 0.8, 1.1, 0.9, 0.9], ["left", "left"] + ["right"] * 6, font_pt=9.5,
                    row_h=0.3)
            d.text(slide, ML, CB - 0.26, LEFT_W, 0.26, "※ 週は月曜始まり。月初・月末の週は日数が少ない", size=8,
                   color=b.col["subtext"])
            b.insight(slide, INS_X, CT, INS_W, CH, "weekly")

        def r2(b, slide, spec):
            d = b.deck
            wd = a["weekday"]
            cats = [w["label"] for w in wd]
            hw = (LEFT_W - 0.3) / 2
            d.chart_title(slide, ML, CT, hw, "曜日別 1日あたり消化金額")
            d.chart(slide, "column", ML, CT + 0.3, hw, 2.4, cats, [("1日あたり消化金額", [w["spend_per_day"]["value"] for w in wd])],
                    number_format=b.fmt_cur, legend=False)
            d.chart_title(slide, ML + hw + 0.3, CT, hw, "曜日別 CPA")
            d.chart(slide, "column", ML + hw + 0.3, CT + 0.3, hw, 2.4, cats,
                    [("CPA", [b.val(w["metrics"], "cpa") for w in wd])], number_format=b.fmt_cur, legend=False,
                    colors=[b.col["series"][3]])
            rows = [[w["label"], f"{w['days']}日", w["spend_per_day"]["display"], w["cv_per_day"]["display"],
                     b.disp(w["metrics"], "cpa"), b.disp(w["metrics"], "cvr"), b.disp(w["metrics"], "roas")] for w in wd]
            d.table(slide, ML, CT + 2.95, LEFT_W, ["曜日", "日数", "1日あたり消化金額", "1日あたりCV", "CPA", "CVR", "ROAS"],
                    rows, [0.7, 0.7, 1.7, 1.3, 1.1, 0.9, 0.9], font_pt=9.5, row_h=0.3)
            b.insight(slide, INS_X, CT, INS_W, CH, "weekday")
        return [Spec("weekly", "週別推移", r1, "weekly", ["total"], "values の week.W<n>.*", slot="weekly"),
                Spec("weekday", "曜日別", r2, "weekly", ["total"],
                     "values の weekday.<mon〜sun>.*（1日あたり = 合計 ÷ その曜日の日数）", slot="weekday")]

    def sec_campaign_compare(self) -> List[Spec]:
        a = self.a

        def r1(b, slide, spec):
            d = b.deck
            rows, styles = [], []
            for c in a["campaigns"]:
                m = c["metrics"]
                share = b.v.get(f"campaign.{c['id']}.spend")
                sp = C.NA_DISPLAY
                tv = a["totals"]["spend"]["value"]
                if share and tv:
                    sp = f"{share['value'] / tv * 100:.1f}%"
                rows.append([c["name"], c["result_label"], b.disp(m, "spend"), sp, b.disp(m, "impressions"),
                             b.disp(m, "ctr"), b.disp(m, "cpc"), b.disp(m, "cv"), b.disp(m, "cpa"), b.disp(m, "roas"),
                             b.disp(m, "results"), b.disp(m, "cost_per_result")])
            t = a["totals"]
            rows.append(["合計", "—", b.disp(t, "spend"), "100.0%", b.disp(t, "impressions"), b.disp(t, "ctr"),
                         b.disp(t, "cpc"), b.disp(t, "cv"), b.disp(t, "cpa"), b.disp(t, "roas"), C.NA_DISPLAY,
                         C.NA_DISPLAY])
            styles = [{}] * (len(rows) - 1) + [{"bold": True, "fill": b.col["accent_light"]}]
            h = d.table(slide, ML, CT, CW, ["キャンペーン", "結果の種類", "消化金額", "構成比", "IMP", "CTR", "CPC", "CV", "CPA",
                                            "ROAS", "結果", "結果の単価"], rows,
                        [2.2, 1.8, 1.2, 0.8, 1.1, 0.8, 0.8, 0.7, 0.9, 0.8, 1.0, 1.0], ["left", "left"] + ["right"] * 10,
                        font_pt=9.5, row_h=0.34, styles=styles)
            d.text(slide, ML, CT + h + 0.05, CW, 0.3,
                   "※ 構成比は消化金額の構成比。「結果」は目的ごとに意味が違うため合計しない（CV は全キャンペーン共通で「購入」）。",
                   size=8, color=b.col["subtext"])
            y = CT + h + 0.4
            b.insight(slide, ML, y, CW, CB - y, "campaigns")

        def r2(b, slide, spec):
            d = b.deck
            camps = a["campaigns"][:10]
            cats = [short_label(c["name"], 18) for c in camps][::-1]
            hh = (CH - 0.4) / 2
            items = [("spend", "消化金額", b.fmt_cur, 0), ("cpa", "CPA（消化金額 ÷ 購入）", b.fmt_cur, 3),
                     ("roas", "ROAS", "0.00", 4), ("ctr", "CTR（リンク）%", PCT_FMT, 1)]
            hw = (CW - 0.3) / 2
            for i, (k, t, f, ci) in enumerate(items):
                x = ML + (i % 2) * (hw + 0.3)
                y = CT + (i // 2) * (hh + 0.4)
                vals = [b.val(c["metrics"], k) for c in camps][::-1]
                if f == PCT_FMT:
                    vals = [pct100(v) for v in vals]
                d.chart_title(slide, x, y, hw, t)
                d.chart(slide, "bar", x, y + 0.3, hw, hh - 0.32, cats, [(t, vals)], number_format=f, legend=False,
                        labels=True, colors=[b.col["series"][ci]], gap=60)
        return [Spec("campaigns", "キャンペーン比較", r1, "campaign_compare",
                     ["total"] + [f"campaign.{c['id']}" for c in a["campaigns"]], "values の campaign.<ID>.*",
                     slot="campaigns"),
                Spec("campaigns.chart", "キャンペーン比較（グラフ）", r2, "campaign_compare",
                     [f"campaign.{c['id']}" for c in a["campaigns"]], "values の campaign.<ID>.spend / cpa / roas")]

    def sec_campaigns(self) -> List[Spec]:
        specs = []
        cats = self.short_dates()
        for c in self.a["campaigns"]:
            def r1(b, slide, spec, c=c):
                d = b.deck
                info = (f"結果の種類：{c['result_label']}｜リーチ：{c['reach_source']}"
                        + ("｜今月から配信（前月なし）" if c.get("new") else ""))
                d.text(slide, ML, CT - 0.05, LEFT_W, 0.3, info, size=9.5, color=b.col["subtext"])
                b.kpi_cards(slide, ML, CT + 0.3, LEFT_W, 2.8, b.cfg["kpi_cards"]["campaign"], c["metrics"], c["mom"])
                y = CT + 3.3
                d.chart_title(slide, ML, y, LEFT_W, "広告セット別")
                rows, styles = [], []
                for sid in c["adset_ids"]:
                    s = b.adset_by_id[sid]
                    m = s["metrics"]
                    mo = (s.get("mom") or {}).get("cpa")
                    rows.append([s["name"], b.disp(m, "spend"), b.disp(m, "cv"), b.disp(m, "cpa"), b.disp(m, "ctr"),
                                 b.disp(m, "roas"), mo["pct"]["display"] if mo else C.NA_DISPLAY])
                    styles.append({"cells": {6: b.trend_color(mo["trend"]["display"]) if mo else b.col["subtext"]}})
                d.table(slide, ML, y + 0.32, LEFT_W, ["広告セット", "消化金額", "CV", "CPA", "CTR", "ROAS", "CPA前月比"], rows,
                        [2.8, 1.3, 0.8, 1.1, 0.9, 0.9, 1.1], font_pt=9.5, row_h=0.3, styles=styles)
                b.insight(slide, INS_X, CT, INS_W, CH, f"campaign.{c['id']}")

            def r2(b, slide, spec, c=c):
                d = b.deck
                ser = c["daily"]
                d.chart_title(slide, ML, CT, CW, "日別 消化金額" + b.mlabel())
                d.chart(slide, "column", ML, CT + 0.3, CW, 2.35, cats, [("消化金額", ser["spend"])],
                        number_format=b.fmt_cur, legend=False, gap=40)
                hw = (CW - 0.3) / 2
                d.chart_title(slide, ML, CT + 2.8, hw, f"日別 {b.label('cv')}" + b.mlabel())
                d.chart(slide, "column", ML, CT + 3.1, hw, 2.45, cats, [(b.label("cv"), ser["cv"])],
                        number_format="#,##0", legend=False, gap=40, colors=[b.col["series"][1]])
                d.chart_title(slide, ML + hw + 0.3, CT + 2.8, hw, "日別 CPA（CV 0件の日は空白）" + b.mlabel())
                d.chart(slide, "line", ML + hw + 0.3, CT + 3.1, hw, 2.45, cats, [("CPA", ser["cpa"])],
                        number_format=b.fmt_cur, legend=False, colors=[b.col["series"][3]])
            scope = f"campaign.{c['id']}"
            specs.append(Spec(scope, f"キャンペーン：{c['name']}", r1, "campaigns",
                              [scope] + [f"adset.{s}" for s in c["adset_ids"]],
                              f"values の {scope}.* / adset.<ID>.*", slot=scope))
            specs.append(Spec(f"{scope}.daily", f"キャンペーン：{c['name']}（日別推移）", r2, "campaigns", [scope],
                              f"analysis.json の campaigns[{c['id']}].daily（日別の元の数値から再計算）"))
        return specs

    def sec_adsets(self) -> List[Spec]:
        specs = []
        cats = self.short_dates()
        for s in self.a["adsets"]:
            def r(b, slide, spec, s=s):
                d = b.deck
                d.text(slide, ML, CT - 0.05, LEFT_W, 0.3, f"キャンペーン：{s['campaign_name']}"
                       + ("｜今月から配信（前月なし）" if s.get("new") else ""), size=9.5, color=b.col["subtext"])
                b.kpi_cards(slide, ML, CT + 0.3, LEFT_W, 1.25, b.cfg["kpi_cards"]["adset"], s["metrics"], s["mom"],
                            cols=6, gap=0.12, value_pt=16)
                hw = (LEFT_W - 0.3) / 2
                y = CT + 1.7
                d.chart_title(slide, ML, y, hw, "日別 消化金額" + b.mlabel())
                d.chart(slide, "column", ML, y + 0.28, hw, 1.75, cats, [("消化金額", s["daily"]["spend"])],
                        number_format=b.fmt_cur, legend=False, gap=40, font_pt=8)
                d.chart_title(slide, ML + hw + 0.3, y, hw, f"日別 {b.label('cv')}" + b.mlabel())
                d.chart(slide, "column", ML + hw + 0.3, y + 0.28, hw, 1.75, cats, [(b.label("cv"), s["daily"]["cv"])],
                        number_format="#,##0", legend=False, gap=40, font_pt=8, colors=[b.col["series"][1]])
                y = CT + 3.85
                d.chart_title(slide, ML, y, LEFT_W, "広告別")
                rows = []
                for aid in s["ad_ids"]:
                    ad = b.ad_by_id[aid]
                    m = ad["metrics"]
                    rows.append([ad["name"], b.disp(m, "spend"), b.disp(m, "cv"), b.disp(m, "cpa"), b.disp(m, "ctr"),
                                 b.disp(m, "roas")])
                d.table(slide, ML, y + 0.3, LEFT_W, ["広告", "消化金額", "CV", "CPA", "CTR", "ROAS"], rows[:5],
                        [3.2, 1.3, 0.8, 1.1, 0.9, 0.9], font_pt=9.5, row_h=0.28)
                b.insight(slide, INS_X, CT, INS_W, CH, f"adset.{s['id']}")
            scope = f"adset.{s['id']}"
            specs.append(Spec(scope, f"広告セット：{s['name']}", r, "adsets", [scope] + [f"ad.{x}" for x in s["ad_ids"]],
                              f"values の {scope}.* / ad.<ID>.*、analysis.json の adsets[{s['id']}].daily", slot=scope))
        return specs

    def selected_ads(self) -> List[dict]:
        ads = self.a["ads"]
        mx = int(self.cfg["ads"].get("max_slides", 40))
        if len(ads) <= mx:
            return ads
        top = set(a["id"] for a in sorted(ads, key=lambda a: -(self.val(a["metrics"], "spend") or 0))[:mx])
        return [a for a in ads if a["id"] in top]

    def sec_ads(self) -> List[Spec]:
        specs = []
        cats = self.short_dates()
        for ad in self.selected_ads():
            def r(b, slide, spec, ad=ad):
                d = b.deck
                period = C.NA_DISPLAY
                if ad["first_date"]:
                    f, l = ad["first_date"], ad["last_date"]
                    period = f"{int(f[5:7])}/{int(f[8:10])}〜{int(l[5:7])}/{int(l[8:10])}（配信 {ad['active_days']}日）"
                d.text(slide, ML, CT - 0.05, LEFT_W, 0.3,
                       f"{ad['campaign_name']} ＞ {ad['adset_name']}｜{period}"
                       + ("｜今月から配信（前月なし）" if ad.get("new") else ""), size=9.5, color=b.col["subtext"])
                b.kpi_cards(slide, ML, CT + 0.3, LEFT_W, 2.55, b.cfg["kpi_cards"]["ad"], ad["metrics"], ad["mom"],
                            value_pt=18)
                hw = (LEFT_W - 0.3) / 2
                y = CT + 3.05
                d.chart_title(slide, ML, y, hw, "日別 消化金額" + b.mlabel())
                d.chart(slide, "column", ML, y + 0.28, hw, CB - y - 0.3, cats, [("消化金額", ad["daily"]["spend"])],
                        number_format=b.fmt_cur, legend=False, gap=40, font_pt=8)
                d.chart_title(slide, ML + hw + 0.3, y, hw, f"日別 {b.label('cv')}" + b.mlabel())
                d.chart(slide, "column", ML + hw + 0.3, y + 0.28, hw, CB - y - 0.3, cats,
                        [(b.label("cv"), ad["daily"]["cv"])], number_format="#,##0", legend=False, gap=40,
                        font_pt=8, colors=[b.col["series"][1]])
                b.insight(slide, INS_X, CT, INS_W, CH, f"ad.{ad['id']}")
            scope = f"ad.{ad['id']}"
            specs.append(Spec(scope, f"広告：{ad['name']}", r, "ads", [scope],
                              f"values の {scope}.*、analysis.json の ads[{ad['id']}].daily", slot=scope))
        return specs

    def sec_ranking(self) -> List[Spec]:
        a = self.a
        n = int(self.cfg["ranking"].get("top_n", 5))
        n2 = int(self.cfg["ranking"].get("spend_top_n", 10))
        rk = a["rankings"]

        def rows_for(ids: List[str], cols: List[str]) -> List[List[str]]:
            out = []
            for i, aid in enumerate(ids, start=1):
                ad = self.ad_by_id[aid]
                out.append([str(i), ad["name"], ad["adset_name"]] + [self.disp(ad["metrics"], k) for k in cols])
            return out

        def r1(b, slide, spec):
            d = b.deck
            hw = (CW - 0.3) / 2
            cols = ["spend", "cv", "cpa", "roas"]
            hdr = ["#", "広告", "広告セット", "消化金額", "CV", "CPA", "ROAS"]
            wd = [0.4, 2.0, 1.9, 1.15, 0.5, 0.95, 0.9]
            al = ["center", "left", "left", "right", "right", "right", "right"]
            d.chart_title(slide, ML, CT, hw, f"CPAが良い広告 TOP{n}")
            h1 = d.table(slide, ML, CT + 0.3, hw, hdr, rows_for(rk["cpa_best"][:n], cols) or [["", "該当なし", "", "", "", "", ""]],
                         wd, al, font_pt=9, row_h=0.3)
            d.chart_title(slide, ML + hw + 0.3, CT, hw, f"CPAが悪い広告 WORST{n}")
            h2 = d.table(slide, ML + hw + 0.3, CT + 0.3, hw, hdr,
                         rows_for(rk["cpa_worst"][:n], cols) or [["", "該当なし", "", "", "", "", ""]], wd, al, font_pt=9,
                         row_h=0.3)
            y = CT + 0.3 + max(h1, h2) + 0.08
            d.text(slide, ML, y, CW, 0.28, f"※ CV が {a['min_cv_for_ranking']}件以上の広告だけを対象（CVが少ない広告のCPAはぶれが大きいため）",
                   size=8, color=b.col["subtext"])
            y += 0.35
            b.insight(slide, ML, y, CW, CB - y, "ranking")

        def r2(b, slide, spec):
            d = b.deck
            hw = (CW - 0.3) / 2
            d.chart_title(slide, ML, CT, hw, f"消化金額 TOP{n2}")
            d.table(slide, ML, CT + 0.3, hw, ["#", "広告", "広告セット", "消化金額", "CV", "CPA"],
                    rows_for(rk["spend"][:n2], ["spend", "cv", "cpa"]), [0.45, 2.2, 2.3, 1.2, 0.55, 1.0],
                    ["center", "left", "left", "right", "right", "right"], font_pt=9, row_h=0.3)
            d.chart_title(slide, ML + hw + 0.3, CT, hw, "消化はあるが CV 0件の広告")
            zr = rows_for(rk["zero_cv"][:n2], ["spend", "impressions", "ctr"])
            d.table(slide, ML + hw + 0.3, CT + 0.3, hw, ["#", "広告", "広告セット", "消化金額", "IMP", "CTR"],
                    zr or [["", "該当なし", "", "", "", ""]], [0.45, 2.2, 2.3, 1.2, 1.0, 0.8],
                    ["center", "left", "left", "right", "right", "right"], font_pt=9, row_h=0.3)
        scopes = sorted({f"ad.{x}" for k in ("cpa_best", "cpa_worst", "spend", "zero_cv") for x in rk[k][:max(n, n2)]})
        return [Spec("ranking", "上位・下位クリエイティブ（CPA）", r1, "ranking", scopes,
                     "analysis.json の rankings / values の rank.<種類>.<順位>.*", slot="ranking"),
                Spec("ranking.spend", "上位・下位クリエイティブ（消化金額・CV 0件）", r2, "ranking", scopes,
                     "analysis.json の rankings.spend / zero_cv")]

    def breakdown_scopes(self, groups: Sequence[str]) -> List[str]:
        out = []
        for g in groups:
            for it in self.a["breakdowns"].get(g, {}).get("items", []):
                out.append(f"{g}.{it['id']}")
        return out

    def sec_age_gender(self) -> List[Spec]:
        bd = self.a["breakdowns"]
        if not bd.get("age_gender", {}).get("available"):
            return []
        items = bd["age_gender"]["items"]
        ages = [it["label"] for it in bd["age"]["items"]]
        genders = [it["label"] for it in bd["gender"]["items"]]

        def matrix(k: str) -> List[Tuple[str, list]]:
            out = []
            for g in genders:
                vals = []
                for a_ in ages:
                    it = next((x for x in items if x["keys"] == [a_, g]), None)
                    vals.append(self.val(it["metrics"], k) if it else None)
                out.append((g, vals))
            return out

        def r1(b, slide, spec):
            d = b.deck
            d.chart_title(slide, ML, CT, LEFT_W, "年齢×性別 消化金額")
            d.chart(slide, "column", ML, CT + 0.3, LEFT_W, 2.45, ages, matrix("spend"), number_format=b.fmt_cur,
                    colors=[b.col["series"][3], b.col["series"][0], b.col["series"][5]], gap=80, overlap=-10)
            d.chart_title(slide, ML, CT + 2.85, LEFT_W, "年齢×性別 CPA（CV 0件は空白）")
            d.chart(slide, "column", ML, CT + 3.15, LEFT_W, 2.4, ages, matrix("cpa"), number_format=b.fmt_cur,
                    colors=[b.col["series"][3], b.col["series"][0], b.col["series"][5]], gap=80, overlap=-10)
            b.insight(slide, INS_X, CT, INS_W, CH, "age_gender")
        rows = []
        for it in items:
            m = it["metrics"]
            rows.append([it["keys"][0], it["keys"][1], self.disp(m, "spend"), it["spend_share"]["display"],
                         self.disp(m, "impressions"), self.disp(m, "ctr"), self.disp(m, "cv"), self.disp(m, "cvr"),
                         self.disp(m, "cpa"), self.disp(m, "roas")])
        scopes = self.breakdown_scopes(["age_gender"])
        specs = [Spec("age_gender", "年齢×性別", r1, "age_gender", scopes,
                      "values の age_gender.<ID>.* / age.<ID>.* / gender.<ID>.*（キャンペーン合算）", slot="age_gender")]
        specs += self.paged_table_specs("age_gender.table", "年齢×性別（明細）", "age_gender",
                                        ["年齢", "性別", "消化金額", "構成比", "IMP", "CTR", "CV", "CVR", "CPA", "ROAS"], rows,
                                        [0.9, 0.8, 1.3, 0.9, 1.3, 0.9, 0.8, 0.9, 1.1, 0.9],
                                        ["left", "left"] + ["right"] * 8,
                                        self.cfg["rows_per_page"].get("breakdown", 18), scopes,
                                        "values の age_gender.<ID>.*", font_pt=9.5,
                                        top_note="※ 全キャンペーンの合算（目的の違うキャンペーンを含む）。CPA は 消化金額 ÷ 購入。")
        return specs

    def _bd_simple(self, kind: str, title: str) -> List[Spec]:
        bd = self.a["breakdowns"].get(kind, {})
        if not bd.get("available"):
            return []
        items = bd["items"]
        top_n = int(self.cfg.get("breakdown_chart_top_n", 10))
        top = items[:top_n]
        scopes = self.breakdown_scopes([kind])

        def r1(b, slide, spec):
            d = b.deck
            if kind == "device":
                hw = (LEFT_W - 0.3) / 2
                d.chart_title(slide, ML, CT, hw, "消化金額の構成比（%）")
                dcats = [short_label(it["label"]) for it in items][::-1]
                d.chart(slide, "bar", ML, CT + 0.3, hw, 2.6, dcats,
                        [("構成比（%）", [pct100(it["spend_share"]["value"]) for it in items][::-1])],
                        number_format=PCT_FMT, legend=False, labels=True, gap=50)
                d.chart_title(slide, ML + hw + 0.3, CT, hw, "CPA（CV 0件は空白）")
                d.chart(slide, "bar", ML + hw + 0.3, CT + 0.3, hw, 2.6, dcats,
                        [("CPA", [b.val(it["metrics"], "cpa") for it in items][::-1])], number_format=b.fmt_cur,
                        legend=False, labels=True, gap=50, colors=[b.col["series"][3]])
                rows = [[it["label"], b.disp(it["metrics"], "spend"), it["spend_share"]["display"],
                         b.disp(it["metrics"], "ctr"), b.disp(it["metrics"], "cv"), b.disp(it["metrics"], "cvr"),
                         b.disp(it["metrics"], "cpa"), b.disp(it["metrics"], "roas")] for it in items]
                d.table(slide, ML, CT + 3.1, LEFT_W, ["デバイス", "消化金額", "構成比", "CTR", "CV", "CVR", "CPA", "ROAS"],
                        rows, [1.9, 1.3, 0.9, 0.8, 0.7, 0.8, 1.0, 0.8], font_pt=9, row_h=0.28)
            else:
                cats = [short_label(it["label"]) for it in top][::-1]
                hw = (LEFT_W - 0.3) / 2
                d.chart_title(slide, ML, CT, hw, f"消化金額（上位{len(top)}）")
                d.chart(slide, "bar", ML, CT + 0.3, hw, CH - 0.35, cats,
                        [("消化金額", [b.val(it["metrics"], "spend") for it in top][::-1])], number_format=b.fmt_cur,
                        legend=False, gap=50)
                d.chart_title(slide, ML + hw + 0.3, CT, hw, "CPA（CV 0件は空白）")
                d.chart(slide, "bar", ML + hw + 0.3, CT + 0.3, hw, CH - 0.35, cats,
                        [("CPA", [b.val(it["metrics"], "cpa") for it in top][::-1])], number_format=b.fmt_cur,
                        legend=False, gap=50, colors=[b.col["series"][3]])
            b.insight(slide, INS_X, CT, INS_W, CH, kind)
        specs = [Spec(kind, title, r1, kind, scopes, f"values の {kind}.<ID>.*（キャンペーン合算）", slot=kind)]
        if kind != "device":
            rows = [[it["label"], self.disp(it["metrics"], "spend"), it["spend_share"]["display"],
                     self.disp(it["metrics"], "impressions"), self.disp(it["metrics"], "ctr"),
                     self.disp(it["metrics"], "cpc"), self.disp(it["metrics"], "cv"), self.disp(it["metrics"], "cvr"),
                     self.disp(it["metrics"], "cpa"), self.disp(it["metrics"], "roas")] for it in items]
            specs += self.paged_table_specs(f"{kind}.table", f"{title}（明細）", kind,
                                            [title.replace("別", ""), "消化金額", "構成比", "IMP", "CTR", "CPC", "CV", "CVR",
                                             "CPA", "ROAS"], rows, [2.8, 1.2, 0.8, 1.2, 0.8, 0.8, 0.7, 0.8, 1.0, 0.8],
                                            None, self.cfg["rows_per_page"].get("breakdown", 18), scopes,
                                            f"values の {kind}.<ID>.*", font_pt=9.5,
                                            top_note="※ 全キャンペーンの合算（目的の違うキャンペーンを含む）。CPA は 消化金額 ÷ 購入。")
        return specs

    def sec_placement(self) -> List[Spec]:
        return self._bd_simple("placement", "配置別")

    def sec_device(self) -> List[Spec]:
        return self._bd_simple("device", "デバイス別")

    def sec_region(self) -> List[Spec]:
        return self._bd_simple("region", "地域別")

    def sec_issues(self) -> List[Spec]:
        def r(b, slide, spec):
            b.insight(slide, ML, CT, CW, CH, "issues", title="今月の所見・課題")
        return [Spec("issues", "所見・課題", r, "issues", ["total"], "insights.json slides.issues（数字はプレースホルダ）",
                     slot="issues")]

    def sec_next_actions(self) -> List[Spec]:
        def r(b, slide, spec):
            d, col = b.deck, b.col
            b.insight(slide, ML, CT, LEFT_W, CH, "next_actions", title="次月施策")
            _, srcs = b.resolved.get("next_actions", ("", []))
            files = [s for s in (srcs or b.slots.get("next_actions", {}).get("sources", [])) if C.is_file_source(s)]
            d.rect(slide, INS_X, CT, INS_W, CH, fill=col["bg_light"])
            written = bool(b.resolved.get("next_actions", ("", []))[0])
            d.text(slide, INS_X + 0.15, CT + 0.1, INS_W - 0.3, 0.3, "根拠にした記録" if written else "参照できる記録（今月）",
                   size=11, bold=True, color=col["primary"])
            lines = ["・" + f.split("/")[-1].rsplit(".", 1)[0] for f in files] or ["今月の議事録・提案書はありません"]
            pt, _ = fit_pt(lines, INS_W - 0.3, CH - 0.6, 10, 7)
            d.text(slide, INS_X + 0.15, CT + 0.5, INS_W - 0.3, CH - 0.6, lines, size=pt, color=col["text"],
                   space_after=3)
        return [Spec("next_actions", "次月施策", r, "next_actions", [],
                     "insights.json slides.next_actions（その月の minutes / proposals を sources に持つ）",
                     slot="next_actions")]

    def sec_appendix(self) -> List[Spec]:
        a = self.a
        specs: List[Spec] = []
        rows, styles = [], []
        for c in a["campaigns"]:
            m = c["metrics"]
            rows.append([c["name"], self.disp(m, "spend"), self.disp(m, "impressions"), self.disp(m, "link_clicks"),
                         self.disp(m, "ctr"), self.disp(m, "cpc"), self.disp(m, "cpm"), self.disp(m, "cv"),
                         self.disp(m, "cvr"), self.disp(m, "cpa"), self.disp(m, "roas")])
            styles.append({"bold": True, "fill": self.col["accent_light"]})
            for sid in c["adset_ids"]:
                s = self.adset_by_id[sid]
                m = s["metrics"]
                rows.append(["　" + s["name"], self.disp(m, "spend"), self.disp(m, "impressions"),
                             self.disp(m, "link_clicks"), self.disp(m, "ctr"), self.disp(m, "cpc"), self.disp(m, "cpm"),
                             self.disp(m, "cv"), self.disp(m, "cvr"), self.disp(m, "cpa"), self.disp(m, "roas")])
                styles.append({})
        hdr = ["キャンペーン／広告セット", "消化金額", "IMP", "リンククリック", "CTR", "CPC", "CPM", "CV", "CVR", "CPA", "ROAS"]
        wd = [3.0, 1.2, 1.1, 1.0, 0.8, 0.8, 0.8, 0.7, 0.8, 0.9, 0.8]
        specs += self.paged_table_specs("appendix.campaigns", "Appendix：キャンペーン・広告セット別 明細", "appendix", hdr,
                                        rows, wd, None, self.cfg["rows_per_page"].get("appendix", 15),
                                        [f"campaign.{c['id']}" for c in a["campaigns"]],
                                        "values の campaign.<ID>.* / adset.<ID>.*", font_pt=9, styles=styles)
        rows = []
        for ad in a["ads"]:
            m = ad["metrics"]
            rows.append([ad["campaign_name"], ad["adset_name"], ad["name"], self.disp(m, "spend"),
                         self.disp(m, "impressions"), self.disp(m, "link_clicks"), self.disp(m, "ctr"),
                         self.disp(m, "cv"), self.disp(m, "cpa"), self.disp(m, "roas")])
        specs += self.paged_table_specs("appendix.ads", "Appendix：広告別 明細", "appendix",
                                        ["キャンペーン", "広告セット", "広告", "消化金額", "IMP", "リンククリック", "CTR", "CV",
                                         "CPA", "ROAS"], rows, [1.9, 2.0, 2.2, 1.1, 1.0, 1.0, 0.7, 0.6, 0.9, 0.7],
                                        ["left", "left", "left"] + ["right"] * 7,
                                        self.cfg["rows_per_page"].get("appendix", 15), ["total"],
                                        "values の ad.<ID>.*（各広告の出どころは広告スライドのノート）", font_pt=9)
        if self.cfg.get("appendix", {}).get("ad_daily"):
            rows = []
            for ad in a["ads"]:
                ser = ad["daily"]
                for i, dd in enumerate(ser["dates"]):
                    if not ser["spend"][i]:
                        continue
                    cpa = ser["cpa"][i]
                    rows.append([dd[5:], ad["adset_name"], ad["name"], f"{ser['spend'][i]:,.0f}",
                                 f"{ser['impressions'][i]:,.0f}", f"{ser['link_clicks'][i]:,.0f}",
                                 f"{ser['cv'][i]:,.0f}" if ser["cv"][i] is not None else C.NA_DISPLAY,
                                 f"{cpa:,.0f}" if cpa is not None else C.NA_DISPLAY])
            specs += self.paged_table_specs("appendix.ad_daily", "Appendix：広告×日 明細", "appendix",
                                            ["日付", "広告セット", "広告", "消化金額", "IMP", "リンククリック", "CV", "CPA"], rows,
                                            [0.8, 2.4, 2.6, 1.2, 1.2, 1.1, 0.8, 1.0], ["left", "left", "left"] + ["right"] * 5,
                                            20, ["total"], "analysis.json の ads[].daily", font_pt=8.5)
        return specs

    def sec_glossary(self) -> List[Spec]:
        a = self.a
        keys = ["spend", "impressions", "reach", "frequency", "link_clicks", "clicks_all", "ctr", "ctr_all", "cpc",
                "cpm", "cv", "cvr", "cpa", "purchase_value", "roas", "aov", "results", "cost_per_result"]
        rows = [[self.defs[k]["label"], self.defs[k]["formula"]] for k in keys if k in self.defs]
        rows.append(["前月比（増減率）", "（今月 − 前月）÷ 前月。比率の指標の差分は pt（ポイント）で表示"])
        rows.append(["※（参考値）", "足し算できない値（リーチ等）を日別・内訳で合算したもの。重複を含む"])
        return self.paged_table_specs("glossary", "Appendix：用語定義", "glossary", ["指標", "定義（計算式）"], rows,
                                      [1.6, 6.0], ["left", "left"], 16, [],
                                      "analysis.json の definitions", font_pt=9.5)

    # ---------------------------------------------------------------- 組み立て
    def plan(self) -> List[Spec]:
        specs: List[Spec] = []
        groups = self.cfg.get("groups", {})
        for sec in self.cfg["sections"]:
            if not sec.get("enabled", True):
                continue
            fn = getattr(self, f"sec_{sec['id']}", None)
            if fn is None:
                self.deck.warnings.append(f"deck_config.json の sections に不明な id があります: {sec['id']}")
                continue
            got = fn()
            if not got:
                self.deck.warnings.append(f"セクション「{SECTION_TITLES.get(sec['id'], sec['id'])}」はデータが無いため省略しました")
            for s in got:
                s.group = sec.get("group", "")
                s.group_label = groups.get(s.group, s.group) if s.group else ""
                if sec.get("title") and s.id == sec["id"]:
                    s.title = sec["title"]
            specs += got
        # 中扉
        if self.cfg.get("section_dividers", True):
            out: List[Spec] = []
            last = None
            for s in specs:
                if s.group and s.group != last:
                    g, gl = s.group, s.group_label
                    members = []
                    for t in specs:
                        if t.group == g and SECTION_TITLES.get(t.section, t.section) not in members:
                            members.append(SECTION_TITLES.get(t.section, t.section))

                    def rdiv(b, slide, spec, gl=gl, members=members):
                        d, col = b.deck, b.col
                        d.rect(slide, 0, 0, W_IN, H_IN, fill=col["primary"])
                        d.rect(slide, 0.9, 2.6, 0.12, 1.9, fill=col["accent"])
                        d.text(slide, 1.25, 2.4, 11, 0.5, spec.extra_notes[0] if spec.extra_notes else "", size=16,
                               color="C9D6EA")
                        d.text(slide, 1.25, 2.85, 11, 1.0, gl, size=36, bold=True, color="FFFFFF")
                        d.text(slide, 1.25, 3.95, 11, 0.6, "　/　".join(members), size=13, color="C9D6EA")
                    div = Spec(f"divider.{g}", gl, rdiv, "divider", kind="divider")
                    div.group, div.group_label = g, gl
                    out.append(div)
                    last = g
                out.append(s)
            specs = out
            n = 0
            for s in specs:
                if s.kind == "divider":
                    n += 1
                    s.extra_notes = ["Appendix" if s.group == "appendix" else f"{n:02d}"]
        for i, s in enumerate(specs, start=1):
            s.page = i
        self.specs = specs
        return specs

    def notes_text(self, spec: Spec) -> str:
        lines = [f"slide_id: {spec.id}", NOTES_MARK]
        if spec.kind in ("cover", "toc", "divider"):
            lines.append("- このスライドに集計値はありません（表紙・目次・中扉）")
            if spec.kind == "cover":
                lines.append(f"- 対象期間: analysis.json の period（{self.a['period']['label']}）")
        else:
            lines.append(f"- 集計ファイル: {self.a['client']}/{self.a['month']}/ads/analysis.json"
                         f"（ad_analyze.py {self.a['generated_at']} 作成）")
            if spec.keys:
                lines.append(f"- 参照: {spec.keys}")
            merged: Dict[str, set] = {}
            for sc in spec.scopes:
                for p in self.a["provenance"].get(sc, []):
                    if "rows" in p:
                        merged.setdefault(p["file"], set()).update(parse_rows(p["rows"]))
                    elif p.get("file"):
                        merged.setdefault(p["file"] + "（" + p.get("field", "") + "）", set())
            for f, rows in merged.items():
                f = notes_path(f)
                if rows:
                    lines.append(f"- 元データ: {f} 行 {compress(rows)}（{len(rows)}行）")
                else:
                    lines.append(f"- 元データ: {f}")
            if spec.scopes:
                lines.append(f"- provenance のスコープ: {', '.join(spec.scopes[:12])}"
                             + (f" ほか{len(spec.scopes) - 12}件" if len(spec.scopes) > 12 else ""))
            lines.append("- 比率（CTR・CPC・CPM・CVR・CPA・ROAS）は元の数値から再計算。CVR の分母はリンククリック。"
                         "CTR はリンクCTR（CTR（すべて）とは別）。定義は Appendix「用語定義」。")
        lines += spec.extra_notes if spec.kind != "divider" else []
        if spec.slot:
            text, srcs = self.resolved.get(spec.slot, ("", []))
            lines.append("【所見の出どころ】")
            if text:
                lines.append(f"- insights.json slides[\"{spec.slot}\"]")
                lines.append("- sources: " + (" / ".join(srcs) if srcs else "（なし）"))
            else:
                lines.append(f"- {INSIGHT_EMPTY_NOTE}（slides[\"{spec.slot}\"] が空）→ スライドは「所見：要確認」")
        return "\n".join(lines)

    def build(self) -> None:
        footer = self.cfg.get("footer", "{client} 様｜{month_label} 広告運用レポート").format(
            client=self.a["client"], month_label=self.a["period"]["month_label"], month=self.a["month"])
        for spec in self.specs:
            slide = self.deck.new_slide()
            if spec.kind in ("content", "toc"):
                self.deck.frame(slide, spec, spec.page, footer)
            spec.render(self, slide, spec)
            slide.notes_slide.notes_text_frame.text = self.notes_text(spec)
        latin, ea = self.deck.font_latin, self.deck.font_ea
        for slide in self.deck.prs.slides:
            apply_fonts_everywhere(slide._element, latin, ea)
            for shp in slide.shapes:
                if getattr(shp, "has_chart", False) and shp.has_chart:
                    apply_fonts_everywhere(shp.chart._chartSpace, latin, ea)
            if slide.has_notes_slide:
                apply_fonts_everywhere(slide.notes_slide._element, latin, ea)


SECTION_TITLES = {
    "cover": "表紙", "toc": "目次", "summary": "エグゼクティブサマリー", "targets": "目標対比", "mom": "前月比較",
    "measurement_notes": "計測上の注意", "daily": "日別推移", "weekly": "週別・曜日別", "campaign_compare": "キャンペーン比較",
    "campaigns": "キャンペーン別", "adsets": "広告セット別", "ads": "広告別", "ranking": "上位・下位", "age_gender": "年齢×性別",
    "placement": "配置", "device": "デバイス", "region": "地域", "issues": "所見・課題", "next_actions": "次月施策",
    "appendix": "明細", "glossary": "用語定義", "divider": "中扉",
}


def notes_path(p: str) -> str:
    """ノートに書くファイルパス。vault の外（絶対パス）は末尾3階層だけにする（ノートは先方にも渡るため）。"""
    if p.startswith("/") or (len(p) > 2 and p[1] == ":"):
        parts = [x for x in p.replace("\\", "/").split("/") if x]
        return ".../" + "/".join(parts[-3:])
    return p


def parse_rows(s: str) -> set:
    out = set()
    for part in str(s).split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            out.update(range(int(a), int(b) + 1))
        else:
            out.add(int(part))
    return out


def compress(rows: set) -> str:
    nums = sorted(rows)
    parts = []
    start = prev = nums[0]
    for n in nums[1:]:
        if n == prev + 1:
            prev = n
            continue
        parts.append(f"{start}-{prev}" if start != prev else str(start))
        start = prev = n
    parts.append(f"{start}-{prev}" if start != prev else str(start))
    return ",".join(parts)


# ---------------------------------------------------------------- CLI

def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="analysis.json と insights.json から広告運用レポートの pptx を作る")
    ap.add_argument("--client", required=True)
    ap.add_argument("--month", help="対象月 YYYY-MM（省略時は前月）")
    ap.add_argument("--template", help="会社テンプレートの .pptx（スライドマスター/レイアウトを使う）")
    ap.add_argument("--config", help="deck_config.json の上書き設定")
    ap.add_argument("--vault", help="vault ルート（省略時はこのスクリプトから自動判定）")
    ap.add_argument("--insights", help="insights.json の場所（既定: analysis.json と同じフォルダ）")
    ap.add_argument("--out", help="出力する pptx のパス（既定: <client>/<YYYY-MM>/ads/<client>_<YYYY-MM>_広告運用レポート.pptx）")
    args = ap.parse_args(argv)

    month = args.month or C.previous_month_of_today()
    if not C.MONTH_RE.match(month):
        print(f"エラー: 月は YYYY-MM 形式で指定してください（受け取った値: {month}）", file=sys.stderr)
        return 2
    vault = C.resolve_vault(args.vault)
    client_dir = C.find_client_dir(vault, args.client)
    if client_dir is None:
        print(f"エラー: クライアント「{args.client}」のフォルダが {vault / '01_clients'} にありません", file=sys.stderr)
        return 2
    client = C.nfc(client_dir.name)
    ads_dir = client_dir / month / "ads"
    apath = ads_dir / "analysis.json"
    if not apath.is_file():
        print(f"エラー: {apath} がありません。先に ad_analyze.py を実行してください", file=sys.stderr)
        return 2
    analysis = C.load_json(apath)
    ipath = Path(args.insights) if args.insights else ads_dir / "insights.json"
    insights: dict = {}
    if ipath.is_file():
        try:
            insights = C.load_json(ipath)
        except ValueError as e:
            print(f"エラー: {ipath} が JSON として読めません: {e}", file=sys.stderr)
            return 1
    try:
        cfg = load_config(args.config)
    except (OSError, ValueError) as e:
        print(f"エラー: 設定ファイルを読めません: {e}", file=sys.stderr)
        return 2
    template = Path(args.template) if args.template else None
    if template and not template.is_file():
        print(f"エラー: テンプレートがありません: {template}", file=sys.stderr)
        return 2

    deck = Deck(cfg, template)
    b = Builder(analysis, insights, cfg, deck, client_dir, vault, ipath if ipath.is_file() else None)
    errors = b.resolve_insights()
    if errors:
        print("エラー: insights.json に解決できないプレースホルダがあります（pptx は作っていません）:", file=sys.stderr)
        for e in errors:
            print(f"  - {e}", file=sys.stderr)
        print("  → analysis.json の values にあるキーだけが使えます（insights.template.json の placeholders を参照）",
              file=sys.stderr)
        C.append_run_log(vault, SCRIPT_NAME, {"client": client, "month": month, "error": "unknown_placeholders",
                                              "count": len(errors)})
        return 1
    specs = b.plan()
    b.build()
    out = Path(args.out) if args.out else ads_dir / f"{client}_{month}_広告運用レポート.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    deck.prs.save(str(out))

    by_sec: Dict[str, int] = {}
    for s in specs:
        k = SECTION_TITLES.get(s.section, s.section)
        by_sec[k] = by_sec.get(k, 0) + 1
    print(f"build_deck: {client} {month}")
    print(f"  insights.json: {'あり' if ipath.is_file() else 'なし（全所見が「要確認」）'}")
    print(f"  スライド {len(specs)}枚 / 所見：要確認 {b.unknown_count}件")
    print("  内訳: " + "、".join(f"{k} {v}" for k, v in by_sec.items()))
    for w in deck.warnings:
        print(f"  [WARN] {w}", file=sys.stderr)
    print(f"  → {C.rel_to(out, vault)}")
    C.append_run_log(vault, SCRIPT_NAME, {"client": client, "month": month, "slides": len(specs),
                                          "unknown_insights": b.unknown_count, "warnings": len(deck.warnings),
                                          "template": str(template) if template else None,
                                          "path": C.rel_to(out, vault)})
    return 0


if __name__ == "__main__":
    sys.exit(main())
